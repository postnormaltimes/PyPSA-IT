from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


WORK_DIR = Path(__file__).resolve().parent
PHASE_ROOT = WORK_DIR.parent
HB_ROOT = PHASE_ROOT / "historical_baseline"
NORMALIZED = HB_ROOT / "normalized"
ANALYSIS = HB_ROOT / "analysis"
QA = PHASE_ROOT / "qa"
RELEASE_ROOT = PHASE_ROOT.parent
ZONES = {"NORD", "CNOR", "CSUD", "SUD", "CALA", "SICI", "SARD"}


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def number(value) -> float:
    return 0.0 if value in (None, "") else float(value)


def close(actual: float, expected: float, tolerance: float = 1e-6) -> bool:
    return abs(actual - expected) <= tolerance


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


checks = []


def check(name: str, passed: bool, actual=None, expected=None, severity="HARD"):
    checks.append({"check": name, "passed": bool(passed), "actual": actual, "expected": expected, "severity": severity})


capacity_path = NORMALIZED / "MEM_Historical_Capacity_By_Zone_Technology.csv"
generation_path = NORMALIZED / "MEM_Historical_Generation_By_Zone_Technology.csv"
storage_path = NORMALIZED / "MEM_Historical_Storage_By_Zone_Technology.csv"
taxonomy_path = NORMALIZED / "MEM_HISTORICAL_TECHNOLOGY_TAXONOMY.csv"
taxonomy_qa_path = HB_ROOT / "qa" / "MEM_HISTORICAL_TAXONOMY_INTEGRITY_QA.csv"
hydro_audit_path = ANALYSIS / "MEM_Hydro_Physical_Inventory_Type_Audit.csv"
matrix_csv_path = ANALYSIS / "MEM_2024_Installed_Capacity_By_Zone_Detailed.csv"
matrix_xlsx_path = ANALYSIS / "MEM_2024_Installed_Capacity_By_Zone_Detailed.xlsx"

capacity = read_csv(capacity_path)
generation = read_csv(generation_path)
storage = read_csv(storage_path)
taxonomy = read_csv(taxonomy_path)
taxonomy_qa = read_csv(taxonomy_qa_path)
hydro_audit = read_csv(hydro_audit_path)
matrix = read_csv(matrix_csv_path)

check("capacity_row_count", len(capacity) == 758, len(capacity), 758)
check("generation_row_count", len(generation) == 881, len(generation), 881)
check("storage_row_count", len(storage) == 19, len(storage), 19)
check("taxonomy_row_count", len(taxonomy) == 40, len(taxonomy), 40)
check("taxonomy_integrity_group_count", len(taxonomy_qa) == 47, len(taxonomy_qa), 47)
taxonomy_failures = [row for row in taxonomy_qa if str(row.get("status", "")).upper().startswith("FAIL")]
check("taxonomy_integrity_zero_failures", not taxonomy_failures, len(taxonomy_failures), 0)

class_counts = Counter(row["hydro_class"] for row in hydro_audit)
check("hydro_physical_inventory_row_count", len(hydro_audit) == 196, len(hydro_audit), 196)
check("hydro_class_counts", class_counts == Counter({"HDAM": 173, "HPHS": 23}), dict(class_counts), {"HDAM": 173, "HPHS": 23})
hphs_rows = [row for row in hydro_audit if row["hydro_class"] == "HPHS"]
hdam_rows = [row for row in hydro_audit if row["hydro_class"] == "HDAM"]
hphs_charge = sum(number(row["charge_power_MW"]) for row in hphs_rows)
hphs_energy = sum(number(row["physical_inventory_energy_MWh"]) for row in hphs_rows)
hdam_energy = sum(number(row["physical_inventory_energy_MWh"]) for row in hdam_rows)
check("HPHS_charge_power_MW", close(hphs_charge, 6809.3), hphs_charge, 6809.3)
check("HPHS_physical_inventory_energy_MWh", close(hphs_energy, 626262.056948, 1e-5), hphs_energy, 626262.056948)
check("HDAM_physical_inventory_energy_MWh", close(hdam_energy, 5748189.091752, 1e-4), hdam_energy, 5748189.091752)

phs_2024_zone = [row for row in storage if row["year"] == "2024" and row["technology"] == "PUMPED_HYDRO" and row["market_zone"] in ZONES]
phs_discharge = sum(number(row["discharge_power_MW"]) for row in phs_2024_zone)
phs_charge = sum(number(row["charge_power_MW"]) for row in phs_2024_zone)
phs_energy = sum(number(row["energy_capacity_MWh"]) for row in phs_2024_zone)
check("canonical_2024_PHS_discharge_MW", close(phs_discharge, 7252.3, 1e-5), phs_discharge, 7252.3)
check("canonical_2024_PHS_charge_MW", close(phs_charge, 6809.3, 1e-6), phs_charge, 6809.3)
check("canonical_2024_PHS_energy_MWh", close(phs_energy, 53000.0, 1e-6), phs_energy, 53000.0)
check("PHS_roundtrip_efficiency_not_invented", all(row["roundtrip_efficiency"] == "" for row in phs_2024_zone), [row["roundtrip_efficiency"] for row in phs_2024_zone], "all blank")

matrix_by_tech = {row["Technology"]: row for row in matrix}
check("matrix_thermoelectric_total_GW", close(number(matrix_by_tech["THERMOELECTRIC_CONVERSION_TOTAL"]["Italy"]), 60.33037427, 1e-9), number(matrix_by_tech["THERMOELECTRIC_CONVERSION_TOTAL"]["Italy"]), 60.33037427)
check("matrix_geothermal_GW", close(number(matrix_by_tech["GEOTHERMAL"]["Italy"]), 0.77179, 1e-9), number(matrix_by_tech["GEOTHERMAL"]["Italy"]), 0.77179)
check("DDS_comparable_thermal_plus_geothermal_GW", close(number(matrix_by_tech["THERMOELECTRIC_CONVERSION_TOTAL"]["Italy"]) + number(matrix_by_tech["GEOTHERMAL"]["Italy"]), 61.10216427, 1e-9), number(matrix_by_tech["THERMOELECTRIC_CONVERSION_TOTAL"]["Italy"]) + number(matrix_by_tech["GEOTHERMAL"]["Italy"]), 61.10216427)
check("matrix_total_hydro_control_GW", close(number(matrix_by_tech["TOTAL_HYDRO_CONTROL"]["Italy"]), 23.294, 1e-8), number(matrix_by_tech["TOTAL_HYDRO_CONTROL"]["Italy"]), 23.294)
check("matrix_bioenergy_cross_classification_GW", close(number(matrix_by_tech["BIOENERGY_RENEWABLE_SOURCE_CROSS_CLASSIFICATION"]["Italy"]), 3.8000923, 1e-9), number(matrix_by_tech["BIOENERGY_RENEWABLE_SOURCE_CROSS_CLASSIFICATION"]["Italy"]), 3.8000923)
check("matrix_bioenergy_non_additive", matrix_by_tech["BIOENERGY_RENEWABLE_SOURCE_CROSS_CLASSIFICATION"]["Accounting role"] == "CROSS_CLASSIFICATION_NON_ADDITIVE", matrix_by_tech["BIOENERGY_RENEWABLE_SOURCE_CROSS_CLASSIFICATION"]["Accounting role"], "CROSS_CLASSIFICATION_NON_ADDITIVE")

wb_formula = load_workbook(matrix_xlsx_path, data_only=False)
wb_values = load_workbook(matrix_xlsx_path, data_only=True)
check("analytical_workbook_sheet_names", wb_formula.sheetnames == ["2024 Capacity by Zone", "2024 Capacity Reconciliation"], wb_formula.sheetnames, ["2024 Capacity by Zone", "2024 Capacity Reconciliation"])
formula_count = sum(1 for ws in wb_formula.worksheets for row in ws.iter_rows() for cell in row if cell.data_type == "f")
formula_errors = []
for ws in wb_values.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("#"):
                formula_errors.append({"sheet": ws.title, "cell": cell.coordinate, "value": cell.value})
check("analytical_workbook_formulas_present", formula_count > 0, formula_count, ">0")
check("analytical_workbook_zero_formula_errors", not formula_errors, formula_errors, [])

provenance = json.loads((QA / "MEM_v2.9_ACCEPTED_EVIDENCE_SEMANTIC_COMPARISON.json").read_text(encoding="utf-8"))
check("v2.9_expected_binary_absent", provenance["binary_match"] is False, provenance["binary_match"], False)
check("v2.9_table_values_match_prechange_evidence", provenance["table_values"]["difference_count"] == 0, provenance["table_values"]["difference_count"], 0)
check("v2.9_formula_anchors_match_prechange_evidence", provenance["formulas_direct_ooxml_controlling"]["difference_count"] == 0, provenance["formulas_direct_ooxml_controlling"]["difference_count"], 0)
check("v2.9_formatting_differences_detected", provenance["computed_styles_normalized_for_equivalent_solid_fill_serializations"]["difference_count"] == 264, provenance["computed_styles_normalized_for_equivalent_solid_fill_serializations"]["difference_count"], 264)

v291_files = list(RELEASE_ROOT.rglob("*v2.9.1*HISTORICAL*INTEGRATED*.xlsx"))
check("no_premature_v2.9.1_workbook", not v291_files, [str(path) for path in v291_files], [])
gem_candidates = [path for path in RELEASE_ROOT.rglob("*.xlsx") if "gem" in path.name.lower() or "gogpt" in path.name.lower()]
check("complete_GEM_Aug2026_workbook_absent", not gem_candidates, [str(path) for path in gem_candidates], [], severity="GATE")

hard_failures = [item for item in checks if item["severity"] == "HARD" and not item["passed"]]
gate_failures = [item for item in checks if item["severity"] == "GATE" and not item["passed"]]
result = {
    "generated_at": datetime.now(timezone.utc).isoformat(),
    "verification_status": "PASS_WITH_EXTERNAL_GATES" if not hard_failures else "FAIL",
    "hard_failure_count": len(hard_failures),
    "gate_failure_count": len(gate_failures),
    "checks": checks,
    "key_results": {
        "thermoelectric_NET_MW": 60330.37427,
        "DDS_comparable_thermal_plus_geothermal_MW": 61102.16427,
        "total_hydro_NET_MW": 23294.0,
        "pure_pump_MW": 3969.57561,
        "mixed_pump_MW": 3282.72439,
        "pumped_discharge_MW": phs_discharge,
        "HPHS_charge_MW": hphs_charge,
        "HPHS_physical_inventory_energy_MWh": hphs_energy,
        "canonical_Terna_PHS_energy_MWh_approx": phs_energy,
        "bioenergy_renewable_source_MW": 3800.0923,
        "bioenergy_thermoelectric_fuel_MW": 3586.0,
        "capacity_rows": len(capacity),
        "generation_rows": len(generation),
        "storage_rows": len(storage),
        "taxonomy_integrity_groups": len(taxonomy_qa),
        "taxonomy_failures": len(taxonomy_failures),
        "analytical_workbook_sha256": sha256(matrix_xlsx_path),
    },
    "gates": {
        "HISTORICAL_EMPIRICAL_BASELINE_STATUS": "INCOMPLETE",
        "CURRENT_2026_PHYSICAL_FLEET_STATUS": "BLOCKED",
        "WORKBOOK_PROMOTION_STATUS": "BLOCKED",
        "GEM_STATUS": "GEM_AUG2026_REQUIRES_USER_DOWNLOAD",
        "v2.9_provenance": provenance["provenance_decision"],
    },
}
output = QA / "MEM_HISTORICAL_BASELINE_PHASE2_5_FINAL_VERIFICATION.json"
output.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
print(json.dumps({
    "verification_status": result["verification_status"],
    "hard_failure_count": result["hard_failure_count"],
    "gate_failure_count": result["gate_failure_count"],
    "key_results": result["key_results"],
    "gates": result["gates"],
    "output": str(output),
}, indent=2))

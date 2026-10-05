from __future__ import annotations

import hashlib
import json
import os
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


WORK_DIR = Path(__file__).resolve().parent
PHASE_ROOT = WORK_DIR.parent
RELEASE_ROOT = PHASE_ROOT.parent
CURRENT = RELEASE_ROOT / "Fase_5_PyPSA_Eur_Italia_2040_2050_Model_Ready_v2.9_MEM_ARCHITECTURE_REFINED.xlsx"
EVIDENCE = Path(f"{CURRENT}.inspect.ndjson")
OUTPUT = PHASE_ROOT / "qa" / "MEM_v2.9_CURRENT_OOXML_SEMANTIC_INVENTORY.json"
EXPECTED_SHA = "4a59c80bdf4da33d825ff94ea5617c03c325ca6fbc26233e7ed633ed6f910640"


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def json_safe(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime,)):
        return value.isoformat()
    if isinstance(value, (list, tuple, set)):
        return [json_safe(item) for item in value]
    if isinstance(value, dict):
        return {str(key): json_safe(item) for key, item in value.items()}
    return str(value)


def read_evidence() -> list[dict]:
    records = []
    with EVIDENCE.open("r", encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                records.append(json.loads(line))
    return records


def range_dimensions(ws):
    return ws.calculate_dimension()


def collect_workbook_semantics(path: Path) -> dict:
    wb = load_workbook(path, data_only=False, read_only=False)
    wb_values = load_workbook(path, data_only=True, read_only=False)
    formula_cells = []
    sheets = []
    for ws in wb.worksheets:
        formula_count = 0
        comments = []
        hyperlinks = []
        hidden_rows = []
        hidden_columns = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formula_count += 1
                    formula_cells.append({"sheet": ws.title, "cell": cell.coordinate, "formula": cell.value})
                if cell.comment is not None:
                    comments.append({"cell": cell.coordinate, "text": cell.comment.text, "author": cell.comment.author})
                if cell.hyperlink is not None:
                    hyperlinks.append({"cell": cell.coordinate, "target": cell.hyperlink.target, "location": cell.hyperlink.location})
        for idx, dim in ws.row_dimensions.items():
            if dim.hidden:
                hidden_rows.append(int(idx))
        for key, dim in ws.column_dimensions.items():
            if dim.hidden:
                hidden_columns.append(key)
        data_validations = []
        if ws.data_validations is not None:
            for validation in ws.data_validations.dataValidation:
                data_validations.append({
                    "sqref": str(validation.sqref),
                    "type": validation.type,
                    "operator": validation.operator,
                    "formula1": validation.formula1,
                    "formula2": validation.formula2,
                    "allow_blank": validation.allow_blank,
                })
        conditional_formatting = []
        for conditional_range in ws.conditional_formatting:
            rules = []
            for rule in ws.conditional_formatting[conditional_range]:
                rules.append({
                    "type": rule.type,
                    "priority": rule.priority,
                    "operator": rule.operator,
                    "text": rule.text,
                    "formula": list(rule.formula or []),
                    "dxf_id": rule.dxfId,
                })
            conditional_formatting.append({"range": str(conditional_range), "rules": rules})
        sheets.append({
            "name": ws.title,
            "state": ws.sheet_state,
            "dimension": range_dimensions(ws),
            "formula_cells": formula_count,
            "merged_ranges": sorted(str(item) for item in ws.merged_cells.ranges),
            "hidden_rows": hidden_rows,
            "hidden_columns": hidden_columns,
            "hyperlinks": hyperlinks,
            "data_validations": data_validations,
            "comments": comments,
            "conditional_formatting": conditional_formatting,
            "tables": sorted(ws.tables.keys()),
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
        })
    defined_names = []
    try:
        for name, defined_name in wb.defined_names.items():
            defined_names.append({
                "name": name,
                "attr_text": defined_name.attr_text,
                "hidden": defined_name.hidden,
                "local_sheet_id": defined_name.localSheetId,
            })
    except Exception as exc:  # pragma: no cover - defensive inventory fallback
        defined_names.append({"inventory_error": repr(exc)})
    return {
        "sheet_order": wb.sheetnames,
        "sheets": sheets,
        "formula_cell_count": len(formula_cells),
        "formula_cells": formula_cells,
        "defined_names": defined_names,
        "calculation": json_safe({
            "calc_mode": getattr(wb.calculation, "calcMode", None),
            "full_calc_on_load": getattr(wb.calculation, "fullCalcOnLoad", None),
            "force_full_calc": getattr(wb.calculation, "forceFullCalc", None),
            "calc_id": getattr(wb.calculation, "calcId", None),
        }),
        "style_table_counts": {
            "cell_styles": len(wb._cell_styles),
            "named_styles": len(wb._named_styles),
            "fonts": len(wb._fonts),
            "fills": len(wb._fills),
            "borders": len(wb._borders),
            "number_formats": len(wb._number_formats),
        },
        "properties": json_safe({
            "creator": wb.properties.creator,
            "last_modified_by": wb.properties.lastModifiedBy,
            "created": wb.properties.created,
            "modified": wb.properties.modified,
            "title": wb.properties.title,
            "subject": wb.properties.subject,
            "description": wb.properties.description,
        }),
        "data_only_companion_loaded": bool(wb_values.sheetnames),
    }


def collect_package_inventory(path: Path) -> dict:
    entries = []
    package_times = Counter()
    package_parts = {}
    with zipfile.ZipFile(path, "r") as archive:
        for info in archive.infolist():
            data = archive.read(info.filename)
            stamp = "%04d-%02d-%02dT%02d:%02d:%02d" % info.date_time
            package_times[stamp] += 1
            entries.append({
                "name": info.filename,
                "bytes": info.file_size,
                "compressed_bytes": info.compress_size,
                "crc": f"{info.CRC:08x}",
                "sha256": sha256_bytes(data),
                "zip_timestamp": stamp,
                "compression": info.compress_type,
            })
        for name in [
            "[Content_Types].xml",
            "_rels/.rels",
            "docProps/core.xml",
            "docProps/app.xml",
            "xl/workbook.xml",
            "xl/_rels/workbook.xml.rels",
            "xl/styles.xml",
            "xl/calcChain.xml",
        ]:
            if name in archive.namelist():
                data = archive.read(name)
                package_parts[name] = {
                    "sha256": sha256_bytes(data),
                    "bytes": len(data),
                    "xml_preview": data.decode("utf-8", errors="replace")[:2000] if name.endswith(".xml") else None,
                }
    return {
        "entry_count": len(entries),
        "zip_timestamp_distribution": dict(sorted(package_times.items())),
        "entries": entries,
        "selected_package_parts": package_parts,
    }


evidence = read_evidence()
semantics = collect_workbook_semantics(CURRENT)
formula_map = {(row["sheet"], row["cell"]): row["formula"] for row in semantics["formula_cells"]}
accepted_formula_records = [record for record in evidence if record.get("kind") == "formula"]
formula_comparison = []
for record in accepted_formula_records:
    key = (record["sheet"], record["address"])
    current_formula = formula_map.get(key)
    accepted_formula = record["formula"]
    if current_formula != accepted_formula:
        value_only = load_workbook(CURRENT, data_only=True, read_only=False)[record["sheet"]][record["address"]].value
        non_data_only = load_workbook(CURRENT, data_only=False, read_only=False)[record["sheet"]][record["address"]].value
        formula_comparison.append({
            "sheet": record["sheet"],
            "cell": record["address"],
            "accepted_formula": accepted_formula,
            "current_formula": current_formula,
            "current_cell_value_non_data_only": non_data_only,
            "current_cached_value_data_only": value_only,
        })

evidence_sheet_records = [record for record in evidence if record.get("kind") == "sheet"]
sheet_topology_differences = []
current_by_name = {sheet["name"]: sheet for sheet in semantics["sheets"]}
for record in evidence_sheet_records:
    current_sheet = current_by_name.get(record["name"])
    if current_sheet is None or current_sheet["dimension"] != record["address"] or semantics["sheet_order"][record["index"]] != record["name"]:
        sheet_topology_differences.append({"accepted": record, "current": current_sheet})

package = collect_package_inventory(CURRENT)
result = {
    "generated_at": datetime.now(timezone.utc).isoformat(),
    "current_path": str(CURRENT),
    "current_bytes": CURRENT.stat().st_size,
    "current_sha256": sha256_file(CURRENT),
    "accepted_sha256": EXPECTED_SHA,
    "binary_match": sha256_file(CURRENT) == EXPECTED_SHA,
    "current_mtime_utc": datetime.fromtimestamp(CURRENT.stat().st_mtime, tz=timezone.utc).isoformat(),
    "accepted_evidence_path": str(EVIDENCE),
    "accepted_evidence_mtime_utc": datetime.fromtimestamp(EVIDENCE.stat().st_mtime, tz=timezone.utc).isoformat(),
    "accepted_evidence_counts": dict(Counter(record.get("kind", "UNKNOWN") for record in evidence)),
    "sheet_topology_difference_count": len(sheet_topology_differences),
    "sheet_topology_difference_sample": sheet_topology_differences[:20],
    "accepted_formula_record_count": len(accepted_formula_records),
    "current_physical_formula_cell_count": semantics["formula_cell_count"],
    "accepted_formula_difference_count": len(formula_comparison),
    "accepted_formula_differences": formula_comparison,
    "formula_difference_interpretation": (
        "CURRENT_WORKBOOK_HAS_LITERAL_CACHED_VALUES_WHERE_PRE_MODIFICATION_EVIDENCE_RECORDED_FORMULAS"
        if formula_comparison and all(item["current_formula"] is None for item in formula_comparison)
        else "NO_ACCEPTED_FORMULA_DIFFERENCES" if not formula_comparison else "MIXED_FORMULA_DIFFERENCES"
    ),
    "workbook_semantics": semantics,
    "package_inventory": package,
}
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
print(json.dumps({
    "current_sha256": result["current_sha256"],
    "current_bytes": result["current_bytes"],
    "sheet_count": len(semantics["sheets"]),
    "sheet_topology_difference_count": result["sheet_topology_difference_count"],
    "current_physical_formula_cell_count": result["current_physical_formula_cell_count"],
    "accepted_formula_difference_count": result["accepted_formula_difference_count"],
    "formula_difference_interpretation": result["formula_difference_interpretation"],
    "zip_timestamp_distribution": package["zip_timestamp_distribution"],
    "output": str(OUTPUT),
}, indent=2))

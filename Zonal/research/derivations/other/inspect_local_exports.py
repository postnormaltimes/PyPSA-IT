from pathlib import Path
import openpyxl

for path in sorted(Path(r"source://downloads").glob("Export-DownloadCenterFile*.xlsx")):
    print(f"FILE {path} {path.stat().st_size}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print("SHEETS", wb.sheetnames)
    for ws in wb.worksheets:
        print("SHEET", ws.title, ws.max_row, ws.max_column)
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 8), values_only=True):
            print(tuple(row[:12]))
    print()

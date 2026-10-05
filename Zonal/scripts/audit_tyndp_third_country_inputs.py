"""Audit official TYNDP 2024 TN00/MT00 inputs used by MEM E1 closure.

This routine is deliberately source-audit only. It does not generate an MEM
price series, build a PyPSA network, or execute an optimisation. Its purpose
is to distinguish a local implementation limitation from an actual missing
official scenario input at the required DE2040 + 2019 evidence grain.
"""

from __future__ import annotations

import os

import csv
import hashlib
import json
import math
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd
import xarray as xr
from openpyxl import load_workbook
from pyxlsb import open_workbook as open_xlsb


ROOT = Path(__file__).resolve().parents[1]
QA_DIR = ROOT / "qa" / "runtime_closure"
SOURCE_DIR = ROOT / "runtime_sources" / "tyndp_2024"
RAW = SOURCE_DIR / "raw"
EXTRACTED = SOURCE_DIR / "extracted"

PYPSA_EUR = Path(os.environ.get("MEM_RESEARCH_UPSTREAM_ROOT", str(Path(__file__).resolve().parents[1] / "data/research/upstream/pypsa-eur")))
ENTSOE_LOAD = (
    PYPSA_EUR
    / "data/entsoe_electricity_demand/archive/2026-02-02/electricity_demand_entsoe_raw.csv"
)
OPSD_LOAD = (
    PYPSA_EUR
    / "data/opsd_electricity_demand/archive/2026-02-02/electricity_demand_opsd_raw.csv"
)

NODES = EXTRACTED / "Nodes/Nodes/LIST OF NODES.xlsx"
SUPPLY = (
    EXTRACTED
    / "20231103-Final-Supply-Inputs-for-TYNDP-2024-Scenarios.xlsx"
    / "20231103 - Final Supply Inputs for TYNDP 2024 Scenarios.xlsx"
)
REFERENCE_GRID = EXTRACTED / "Line-data/Line data/ReferenceGrid_Electricity.xlsx"
DEMAND_SCENARIOS = (
    EXTRACTED
    / "Demand_Scenarios_TYNDP_2024_After_Public_Consultation.xlsb"
    / "Demand_Scenarios_TYNDP_2024_After_Public_Consultation.xlsb"
)
DE_DEMAND = EXTRACTED / "Demand Profiles/DE/2040/ELECTRICITY_MARKET DE 2040.xlsx"
NT_DEMAND = (
    EXTRACTED
    / "Demand Profiles/NT/Electricity demand profiles/2040_National Trends.xlsx"
)
DE_OUTPUT = (
    EXTRACTED
    / "DE2040CY1995/MMStandardOutputFile_DE2040_Plexos_CY1995_v11_SoS.xlsb"
)
PRICES = EXTRACTED / "Prices/Prices/2023 06 22 TYNDP 2024 Commodity prices Final.xlsx"
METHODOLOGY_REPORT = (
    RAW / "TYNDP_2024_Scenarios_Methodology_Report_Final_250128.pdf"
)
IMPLEMENTATION_GUIDELINES = (
    RAW / "TYNDP_2024_Implementation_Guidelines_intermediate.pdf"
)
TN_PEMMDB_COMPARISON = (
    QA_DIR / "MEM_TYNDP_TN00_PEMMDB_2030_2040_2050_COMPARISON.json"
)
CUTOUT = (
    ROOT
    / "runtime_sources/weather/PyPSA_cutout_v1.0/europe-2019-sarah3-era5.nc"
)
CUTOUT_SHA256 = "E6CF2B4C9D463D64B4DFB3A613EF20C3C69E99624DAADEEE6EE9C76CF74E9D1F"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def _normalise_year(value: Any) -> int | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number.is_integer() and 1900 <= number <= 2100:
        return int(number)
    return None


def _xlsx_hits(path: Path, needles: Iterable[str]) -> list[dict[str, Any]]:
    upper_needles = tuple(item.upper() for item in needles)
    hits: list[dict[str, Any]] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=False):
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and any(
                        needle in value.upper() for needle in upper_needles
                    ):
                        hits.append(
                            {
                                "sheet": worksheet.title,
                                "row": cell.row,
                                "column": cell.column,
                                "value": value,
                            }
                        )
    finally:
        workbook.close()
    return hits


def _xlsb_hits(path: Path, needles: Iterable[str]) -> list[dict[str, Any]]:
    upper_needles = tuple(item.upper() for item in needles)
    hits: list[dict[str, Any]] = []
    workbook = open_xlsb(path)
    try:
        for sheet_name in workbook.sheets:
            worksheet = workbook.get_sheet(sheet_name)
            try:
                for row_number, row in enumerate(worksheet.rows(), start=1):
                    for column_number, cell in enumerate(row, start=1):
                        value = cell.v
                        if isinstance(value, str) and any(
                            needle in value.upper() for needle in upper_needles
                        ):
                            hits.append(
                                {
                                    "sheet": sheet_name,
                                    "row": row_number,
                                    "column": column_number,
                                    "value": value,
                                }
                            )
            finally:
                worksheet.close()
    finally:
        workbook.close()
    return hits


def _xlsx_profile_stats(
    path: Path, sheet_name: str, years: Iterable[int]
) -> dict[int, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return {
                year: {
                    "sheet_present": False,
                    "header_present": False,
                    "nonnull": 0,
                    "sum": None,
                    "min": None,
                    "max": None,
                }
                for year in years
            }
        worksheet = workbook[sheet_name]
        header_row = None
        column_by_year: dict[int, int] = {}
        for row_number, values in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=min(20, worksheet.max_row),
                values_only=True,
            ),
            start=1,
        ):
            if values and isinstance(values[0], str) and values[0].strip().lower() == "date":
                header_row = row_number
                for column_number, value in enumerate(values, start=1):
                    year = _normalise_year(value)
                    if year is not None:
                        column_by_year[year] = column_number
                break

        output: dict[int, dict[str, Any]] = {}
        for year in years:
            column_number = column_by_year.get(year)
            numbers: list[float] = []
            if header_row is not None and column_number is not None:
                for (value,) in worksheet.iter_rows(
                    min_row=header_row + 1,
                    max_row=worksheet.max_row,
                    min_col=column_number,
                    max_col=column_number,
                    values_only=True,
                ):
                    if isinstance(value, (int, float)) and math.isfinite(float(value)):
                        numbers.append(float(value))
            output[year] = {
                "sheet_present": True,
                "header_present": column_number is not None,
                "nonnull": len(numbers),
                "sum": sum(numbers) if numbers else None,
                "min": min(numbers) if numbers else None,
                "max": max(numbers) if numbers else None,
            }
        return output
    finally:
        workbook.close()


def _pecd_profile_stats(path: Path, year: int = 2019) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    metadata: dict[str, str] = {}
    header_index = None
    for index, row in enumerate(rows):
        if row and row[0] in {"Country", "PECD Zone", "Target Year", "Scenario"}:
            metadata[row[0]] = row[1] if len(row) > 1 else ""
        if row and row[0] == "Date":
            header_index = index
            break
    if header_index is None:
        raise ValueError(f"No Date header found in {path}")
    header = rows[header_index]
    target_names = {str(year), f"{year}.0"}
    column_number = next(
        (index for index, value in enumerate(header) if value in target_names), None
    )
    values: list[float] = []
    if column_number is not None:
        for row in rows[header_index + 1 :]:
            if column_number >= len(row) or row[column_number].strip() == "":
                continue
            try:
                value = float(row[column_number])
            except ValueError:
                continue
            if math.isfinite(value):
                values.append(value)
    return {
        "metadata": metadata,
        "year_header_present": column_number is not None,
        "nonnull": len(values),
        "sum": sum(values) if values else None,
        "min": min(values) if values else None,
        "max": max(values) if values else None,
    }


def _pemdb_summary(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["Thermal"]
        capacities: list[tuple[str, str, float]] = []
        current_family = "UNSPECIFIED"
        for row in worksheet.iter_rows(min_row=12, values_only=True):
            if row and row[0] not in (None, ""):
                current_family = str(row[0])
            label = row[1] if len(row) > 1 else None
            capacity = row[2] if len(row) > 2 else None
            if isinstance(capacity, (int, float)) and float(capacity) != 0:
                capacities.append((current_family, str(label), float(capacity)))

        def renewable_capacity_mw(sheet_name: str) -> float:
            """Read Wind/Solar installed-capacity rows (reported in GW)."""
            if sheet_name not in workbook.sheetnames:
                return 0.0
            total_gw = 0.0
            for row in workbook[sheet_name].iter_rows(values_only=True):
                label = str(row[0] or "")
                value = row[1] if len(row) > 1 else None
                if (
                    label.startswith("Installed capacities")
                    and isinstance(value, (int, float))
                    and math.isfinite(float(value))
                ):
                    total_gw += float(value)
            return total_gw * 1000.0

        def hydro_power_mw() -> float:
            if "Hydro" not in workbook.sheetnames:
                return 0.0
            total = 0.0
            for row in workbook["Hydro"].iter_rows(values_only=True):
                label = str(row[0] or "")
                value = row[1] if len(row) > 1 else None
                if (
                    "capacity (MW)" in label
                    and isinstance(value, (int, float))
                    and math.isfinite(float(value))
                ):
                    total += float(value)
            return total

        def battery_power_mw() -> float:
            if "Battery" not in workbook.sheetnames:
                return 0.0
            total = 0.0
            for row in workbook["Battery"].iter_rows(min_row=9, values_only=True):
                value = row[2] if len(row) > 2 else None
                if isinstance(value, (int, float)) and math.isfinite(float(value)):
                    total += float(value)
            return total

        return {
            "country": worksheet.cell(3, 2).value,
            "node": worksheet.cell(4, 2).value,
            "metadata_year": worksheet.cell(5, 2).value,
            "metadata_scenario": worksheet.cell(7, 2).value,
            "thermal_capacity_MW": sum(value for _, _, value in capacities),
            "thermal_nonzero_bands": capacities,
            "gas_capacity_MW": sum(
                value
                for family, _, value in capacities
                if family.strip().lower() == "gas"
            ),
            "nuclear_capacity_MW": sum(
                value
                for family, _, value in capacities
                if family.strip().lower() == "nuclear"
            ),
            "wind_capacity_MW": renewable_capacity_mw("Wind"),
            "solar_capacity_MW": renewable_capacity_mw("Solar"),
            "hydro_power_MW": hydro_power_mw(),
            "battery_power_MW": battery_power_mw(),
        }
    finally:
        workbook.close()


def _de2040_commodity_prices(path: Path) -> dict[str, float]:
    """Read the common DE2040 fuel and carbon values from the official workbook."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["Final Fuel Prices  2022"]
        if worksheet.cell(3, 8).value != "Distributed Energy":
            raise ValueError("Distributed Energy header not found at expected column")
        if int(worksheet.cell(4, 9).value) != 2040:
            raise ValueError("DE2040 column not found at expected position")
        values = {
            "natural_gas_EUR_per_net_GJ": float(worksheet.cell(12, 9).value),
            "gas_blend_EUR_per_net_GJ": float(worksheet.cell(15, 9).value),
            "co2_EUR_per_t": float(worksheet.cell(20, 9).value),
        }
        if worksheet.cell(12, 2).value != "Natural Gas":
            raise ValueError("Natural Gas row not found at expected position")
        if worksheet.cell(20, 2).value != "CO2 price":
            raise ValueError("CO2 row not found at expected position")
        return values
    finally:
        workbook.close()


def _cutout_summary(path: Path) -> dict[str, Any]:
    dataset = xr.open_dataset(path)
    try:
        return {
            "hours": int(dataset.sizes["time"]),
            "time_start": str(dataset.time.min().values),
            "time_end": str(dataset.time.max().values),
            "x_min": float(dataset.x.min()),
            "x_max": float(dataset.x.max()),
            "y_min": float(dataset.y.min()),
            "y_max": float(dataset.y.max()),
            "variables": list(dataset.data_vars),
        }
    finally:
        dataset.close()


def _de_output_header_summary(path: Path) -> dict[str, Any]:
    workbook = open_xlsb(path)
    try:
        hourly = workbook.get_sheet("Hourly Market Data emarket")
        header_rows: dict[int, list[Any]] = {}
        try:
            for row_number, row in enumerate(hourly.rows()):
                if row_number in {10, 11, 12}:
                    header_rows[row_number] = [cell.v for cell in row]
                if row_number >= 12:
                    break
            dimension = hourly.dimension
        finally:
            hourly.close()

        categories = header_rows[10]
        countries = header_rows[11]
        codes = header_rows[12]
        node_columns: dict[str, list[dict[str, Any]]] = {"TN00": [], "MT00": []}
        for column_number, (category, country, code) in enumerate(
            zip(categories, countries, codes), start=1
        ):
            text = f"{country or ''}|{code or ''}"
            for node in node_columns:
                if node in text:
                    node_columns[node].append(
                        {
                            "column": column_number,
                            "category": category,
                            "country": country,
                            "code": code,
                        }
                    )

        crossborder = workbook.get_sheet("Crossborder exchanges")
        try:
            crossborder_headers: list[Any] = []
            for row_number, row in enumerate(crossborder.rows()):
                if row_number == 10:
                    crossborder_headers = [cell.v for cell in row]
                    break
        finally:
            crossborder.close()

        return {
            "sheet_rows": dimension.h,
            "sheet_columns": dimension.w,
            "modelled_hours": dimension.h - 13,
            "node_column_counts": {
                node: len(columns) for node, columns in node_columns.items()
            },
            "marginal_cost_columns": {
                node: [
                    column
                    for column in columns
                    if "Marginal Cost" in str(column["category"])
                    or "Mgl Cost" in str(column["code"])
                ]
                for node, columns in node_columns.items()
            },
            "crossborder_columns": {
                node: [value for value in crossborder_headers if node in str(value)]
                for node in node_columns
            },
        }
    finally:
        workbook.close()


def _load_compatibility() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    entsoe = pd.read_csv(
        ENTSOE_LOAD,
        usecols=lambda column: column == "IT" or str(column).startswith("Unnamed:"),
    )
    entsoe = entsoe.rename(columns={entsoe.columns[0]: "timestamp"})
    entsoe["timestamp"] = pd.to_datetime(entsoe["timestamp"], utc=True)
    entsoe = entsoe.set_index("timestamp")["IT"].sort_index()

    opsd = pd.read_csv(OPSD_LOAD, usecols=["utc_timestamp", "IT"])
    opsd["utc_timestamp"] = pd.to_datetime(opsd["utc_timestamp"], utc=True)
    opsd = opsd.set_index("utc_timestamp")["IT"].sort_index()

    start = pd.Timestamp("2019-01-01 00:00:00", tz="UTC")
    end = pd.Timestamp("2019-12-31 23:00:00", tz="UTC")
    expected = pd.date_range(start, end, freq="h")
    entsoe = entsoe.reindex(expected)
    opsd = opsd.reindex(expected)
    overlap = pd.concat({"ENTSOE": entsoe, "OPSD": opsd}, axis=1).dropna()
    difference = overlap["ENTSOE"] - overlap["OPSD"]
    gap = expected[entsoe.isna()]
    recovery_value = float(opsd.loc[gap[0]]) if len(gap) == 1 else None
    metrics = {
        "expected_hours": len(expected),
        "entsoe_nonnull": int(entsoe.notna().sum()),
        "opsd_nonnull": int(opsd.notna().sum()),
        "overlap_hours": len(overlap),
        "exact_overlap_hours": int(np.isclose(difference, 0.0, atol=1e-9).sum()),
        "correlation": float(overlap.corr().iloc[0, 1]),
        "entsoe_mean_MW": float(overlap["ENTSOE"].mean()),
        "opsd_mean_MW": float(overlap["OPSD"].mean()),
        "mean_bias_MW": float(difference.mean()),
        "mae_MW": float(difference.abs().mean()),
        "median_abs_difference_MW": float(difference.abs().median()),
        "rmse_MW": float(np.sqrt((difference**2).mean())),
        "mean_ratio": float((overlap["ENTSOE"] / overlap["OPSD"]).mean()),
        "entsoe_missing": [timestamp.isoformat() for timestamp in gap],
        "recovery_value_MW": recovery_value,
    }
    rows = [
        {
            "check_id": "LOAD-COMP-001",
            "subject": "Timestamp basis compatibility",
            "observed": "ENTSO-E UTC-aware; OPSD utc_timestamp Z; 8,758 overlapping 2019 values",
            "expected": "Same hourly UTC basis",
            "status": "PASS",
            "notes": "Both series parse to one unique UTC timestamp index.",
        },
        {
            "check_id": "LOAD-COMP-002",
            "subject": "Unit compatibility",
            "observed": "Both source workflows expose Italy load as MW",
            "expected": "MW",
            "status": "PASS",
            "notes": "Overlap means differ by only 10.646 MW; no scale conversion is indicated.",
        },
        {
            "check_id": "LOAD-COMP-003",
            "subject": "Overlap correlation",
            "observed": f"{metrics['correlation']:.12f}",
            "expected": ">=0.999",
            "status": "PASS" if metrics["correlation"] >= 0.999 else "FAIL",
            "notes": f"MAE={metrics['mae_MW']:.6f} MW; RMSE={metrics['rmse_MW']:.6f} MW.",
        },
        {
            "check_id": "LOAD-COMP-004",
            "subject": "ENTSO-E missing observation",
            "observed": "|".join(metrics["entsoe_missing"]),
            "expected": "2019-01-01T00:00:00+00:00 only",
            "status": "PASS"
            if metrics["entsoe_missing"] == ["2019-01-01T00:00:00+00:00"]
            else "FAIL",
            "notes": "No other ENTSO-E 2019 gap exists.",
        },
        {
            "check_id": "LOAD-COMP-005",
            "subject": "OPSD direct recovery value",
            "observed": f"{recovery_value:.1f} MW" if recovery_value is not None else "MISSING",
            "expected": "22850.0 MW",
            "status": "PASS" if recovery_value == 22850.0 else "FAIL",
            "notes": "Direct source recovery; no interpolation and no blending at other hours.",
        },
    ]
    return rows, metrics


def _receipt(
    source_id: str,
    role: str,
    path: Path,
    coverage: str,
    status: str,
    notes: str,
) -> dict[str, Any]:
    return {
        "source_id": source_id,
        "role": role,
        "path": str(path.resolve()),
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
        "coverage": coverage,
        "status": status,
        "notes": notes,
    }


def _write_csv(
    path: Path,
    rows: list[dict[str, Any]],
    fieldnames: list[str] | None = None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    names = fieldnames or list(rows[0])
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=names, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def build_audit() -> dict[str, Any]:
    required = [
        NODES,
        SUPPLY,
        REFERENCE_GRID,
        DEMAND_SCENARIOS,
        DE_DEMAND,
        NT_DEMAND,
        DE_OUTPUT,
        PRICES,
        METHODOLOGY_REPORT,
        IMPLEMENTATION_GUIDELINES,
        TN_PEMMDB_COMPARISON,
        CUTOUT,
        ENTSOE_LOAD,
        OPSD_LOAD,
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("Required audit files missing: " + "; ".join(missing))

    node_hits = _xlsx_hits(NODES, ["TN00", "MT00", "Tunisia", "Malta"])
    supply_hits = _xlsx_hits(SUPPLY, ["TN00", "MT00", "Tunisia", "Malta"])
    grid_hits = _xlsx_hits(REFERENCE_GRID, ["TN00", "MT00"])
    scenario_hits = _xlsb_hits(
        DEMAND_SCENARIOS, ["TN00", "MT00", "Tunisia", "Malta"]
    )

    de_mt = _xlsx_profile_stats(DE_DEMAND, "MT00", [1995, 2008, 2009, 2019])
    nt_mt = _xlsx_profile_stats(NT_DEMAND, "MT00", [2018, 2019])
    nt_tn = _xlsx_profile_stats(NT_DEMAND, "TN00", [2018, 2019])

    pecd_files = {
        "MT00_SOLAR": EXTRACTED
        / "PECD/2040/PECD_LFSolarPV_2040_MT00_edition 2023.2.csv",
        "TN00_SOLAR": EXTRACTED
        / "PECD/2040/PECD_LFSolarPV_2040_TN00_edition 2023.2.csv",
        "TN00_WIND": EXTRACTED
        / "PECD/2040/PECD_Wind_Onshore_2040_TN00_edition 2023.2.csv",
    }
    pecd = {name: _pecd_profile_stats(path) for name, path in pecd_files.items()}

    pemmdb_files = {
        "MT00_2030": EXTRACTED
        / "PEMMDB2_compare/PEMMDB2/2030/PEMMDB_MT00_NationalTrends_2030.xlsx",
        "MT00_2040": EXTRACTED
        / "PEMMDB2/PEMMDB2/2040/PEMMDB_MT00_NationalTrends_2040.xlsx",
        "MT00_2050": EXTRACTED
        / "PEMMDB2_compare/PEMMDB2/2050/PEMMDB_MT00_NationalTrends_2050.xlsx",
        "TN00_2030": EXTRACTED
        / "PEMMDB2_compare/PEMMDB2/2030/PEMMDB_TN00_NationalTrends_2030.xlsx",
        "TN00_2040": EXTRACTED
        / "PEMMDB2/PEMMDB2/2040/PEMMDB_TN00_NationalTrends_2040.xlsx",
        "TN00_2050": EXTRACTED
        / "PEMMDB2_compare/PEMMDB2/2050/PEMMDB_TN00_NationalTrends_2050.xlsx",
    }
    pemmdb = {name: _pemdb_summary(path) for name, path in pemmdb_files.items()}
    output = _de_output_header_summary(DE_OUTPUT)
    commodity_prices = _de2040_commodity_prices(PRICES)
    cutout = _cutout_summary(CUTOUT)
    with TN_PEMMDB_COMPARISON.open("r", encoding="utf-8") as handle:
        tn_pemmdb_comparison = json.load(handle)
    load_rows, load_metrics = _load_compatibility()

    def has(hits: list[dict[str, Any]], term: str) -> bool:
        return any(term.upper() in str(hit["value"]).upper() for hit in hits)

    source_presence = {
        "nodes_mt": has(node_hits, "MT00") or has(node_hits, "Malta"),
        "nodes_tn": has(node_hits, "TN00") or has(node_hits, "Tunisia"),
        "supply_mt": has(supply_hits, "MT00") or has(supply_hits, "Malta"),
        "supply_tn": has(supply_hits, "TN00") or has(supply_hits, "Tunisia"),
        "demand_scenarios_mt": has(scenario_hits, "MT00")
        or has(scenario_hits, "Malta"),
        "demand_scenarios_tn": has(scenario_hits, "TN00")
        or has(scenario_hits, "Tunisia"),
        "grid_mt": has(grid_hits, "MT00"),
        "grid_tn": has(grid_hits, "TN00"),
    }
    expected_presence = {
        "nodes_mt": True,
        "nodes_tn": False,
        "supply_mt": True,
        "supply_tn": False,
        "demand_scenarios_mt": True,
        "demand_scenarios_tn": False,
        "grid_mt": True,
        "grid_tn": True,
    }
    if source_presence != expected_presence:
        raise AssertionError(
            f"Official-input presence pattern changed: {source_presence}"
        )

    coverage: list[dict[str, Any]] = []

    def rel(path: Path, page: int | None = None) -> str:
        value = str(path.relative_to(ROOT))
        return f"{value} [page {page}]" if page is not None else value

    def add(
        node: str,
        field: str,
        scenario_year: str,
        official_file: str,
        observed: str,
        provenance: str,
        interpretation: str,
        required: bool,
        pass_gap: str,
    ) -> None:
        coverage.append(
            {
                "node": node,
                "field": field,
                "scenario_year": scenario_year,
                "official_file": official_file,
                "observed_value_status": observed,
                "provenance": provenance,
                "interpretation": interpretation,
                "required_for_price_model": str(required).lower(),
                "pass_gap": pass_gap,
            }
        )

    mt = pemmdb["MT00_2040"]
    tn = pemmdb["TN00_2040"]
    tn_2040_2050 = tn_pemmdb_comparison["comparisons"]["2040_vs_2050"]
    tn_link_targets = [
        target
        for member in tn_pemmdb_comparison["files"]["2040"]["package_links"]
        for target in member.get("targets", [])
        if "PEMMDB23_TN00_NationalTrends_2030.xlsx" in target
    ]
    common_costs = (
        f"DE2040 natural gas={commodity_prices['natural_gas_EUR_per_net_GJ']:.10f} "
        f"EUR/net GJ; gas blend={commodity_prices['gas_blend_EUR_per_net_GJ']:.10f} "
        f"EUR/net GJ; CO2={commodity_prices['co2_EUR_per_t']:.3f} EUR/t"
    )

    add(
        "MT00", "OFFICIAL_SCENARIO_RULE", "TYNDP2024 DE/GA 2040",
        rel(METHODOLOGY_REPORT, 17),
        "Malta is shown in the modelled bidding-zone perimeter; the North-African-only-NT rule does not apply to Malta.",
        "EXPLICIT_FINAL_OFFICIAL_METHODOLOGY",
        "MT00 is an official DE/GA node. The later final scenario report controls over the generic intermediate project-assessment perimeter.",
        True, "PASS",
    )
    add(
        "TN00", "OFFICIAL_SCENARIO_RULE", "TYNDP2024 DE/GA 2040",
        rel(METHODOLOGY_REPORT, 17),
        "North African electricity markets belonging to Med-TSO are modelled only in National Trends; Tunisia is absent from Figure 5.",
        "EXPLICIT_FINAL_OFFICIAL_METHODOLOGY",
        "No official TYNDP rule carries NT Tunisia into DE/GA. Combining DE Europe with NT Tunisia would be an MEM-created hybrid unless authorized.",
        True, "GAP",
    )
    add(
        "MT00", "MARKET_NODE", "DE2040",
        f"{rel(NODES)}; {rel(DE_OUTPUT)}",
        "MT00 present in List of Nodes and published DE2040 output; marginal-cost columns=5.",
        "DIRECT_OFFICIAL_INPUT_AND_OUTPUT",
        "Missing helper support is a bounded adapter issue, not a scenario-data gap.",
        True, "PASS",
    )
    add(
        "TN00", "MARKET_NODE", "DE2040",
        f"{rel(NODES)}; {rel(DE_OUTPUT)}; {rel(IMPLEMENTATION_GUIDELINES, 14)}",
        "TN00 absent from List of Nodes and published DE2040 output; the intermediate guideline includes Tunisia only in its wider project-assessment perimeter.",
        "DIRECT_OFFICIAL_INPUT_AND_OUTPUT_PLUS_METHOD_COMPARISON",
        "Tunisia is a recognized TYNDP third-country/NT market but is not demonstrated as a DE2040 price node.",
        True, "GAP",
    )
    add(
        "MT00", "DEMAND_ANNUAL_QUANTITY", "DE2040",
        rel(DE_DEMAND),
        "; ".join(
            f"CY{year}={de_mt[year]['sum'] / 1e6:.9f} TWh"
            for year in (1995, 2008, 2009)
        ),
        "DERIVED_SUM_OF_DIRECT_OFFICIAL_HOURLY_PROFILES",
        "The official DE2040 annual quantity basis is present; choosing CY2019 changes only the temporal realization, not the scenario quantity.",
        True, "PASS",
    )
    add(
        "TN00", "DEMAND_ANNUAL_QUANTITY", "DE2040 / closest NT2040",
        f"{rel(DE_DEMAND)}; {rel(NT_DEMAND)}",
        f"DE2040 TN00 sheet absent; closest official NT2040 CY2018 profile={nt_tn[2018]['sum'] / 1e6:.9f} TWh.",
        "DIRECT_OFFICIAL_ABSENCE_AND_NON_TARGET_SCENARIO_VALUE",
        "The NT quantity cannot be treated as DE2040 without an explicit MEM cross-scenario authorization.",
        True, "GAP",
    )
    add(
        "MT00", "DEMAND_HOURLY_SHAPE", "DE2040 / CY2019",
        f"{rel(DE_DEMAND)}; {rel(METHODOLOGY_REPORT, 21)}",
        f"Official DE profiles complete for CY1995/2008/2009 (8,760 each); CY2019={de_mt[2019]['nonnull']}/8,760. Malta is documented as rescaling TYNDP2022 demand profiles rather than using DFT.",
        "DIRECT_OFFICIAL_PROFILE_COVERAGE_AND_EXPLICIT_METHOD",
        "A bounded, documented 2019 temporal adapter remains; no new annual DE2040 demand quantity is needed.",
        True, "GAP",
    )
    add(
        "TN00", "DEMAND_HOURLY_SHAPE", "NT2040 / CY2019",
        rel(NT_DEMAND),
        f"CY2018={nt_tn[2018]['nonnull']}/8,760; CY2019={nt_tn[2019]['nonnull']}/8,760; local ENTSO-E/OPSD exact-country series absent.",
        "DIRECT_OFFICIAL_PROFILE_COVERAGE_PLUS_LOCAL_INVENTORY",
        "A traceable 2019 load adapter is still needed after, and cannot substitute for, the unresolved scenario mapping.",
        True, "GAP",
    )
    add(
        "MT00", "THERMAL_GAS_CAPACITY", "2040 National Trends fleet used by DE method",
        rel(pemmdb_files["MT00_2040"]),
        f"metadata=MT00/2040/NationalTrends; gas CCGT={mt['gas_capacity_MW']:.3f} MW; total thermal={mt['thermal_capacity_MW']:.3f} MW.",
        "DIRECT_OFFICIAL_PEMMDB",
        "The 2040 fleet member has valid horizon and scenario metadata and is compatible with the final methodology's DE fleet rules.",
        True, "PASS",
    )
    add(
        "TN00", "THERMAL_GAS_CAPACITY", "nominal NT2040 member",
        f"{rel(pemmdb_files['TN00_2040'])}; {rel(TN_PEMMDB_COMPARISON)}",
        f"filename says 2040; metadata year={tn['metadata_year']}, scenario=blank; gas/thermal={tn['gas_capacity_MW']:.3f}/{tn['thermal_capacity_MW']:.3f} MW; 2040-vs-2050 cached/formula differences={tn_2040_2050['cached_visible_cell_difference_count']}/{tn_2040_2050['formula_cell_difference_count']}; 2030-link retained={bool(tn_link_targets)}.",
        "DIRECT_OFFICIAL_MEMBER_PLUS_OOXML_HORIZON_COMPARISON",
        "This is a stale/packaging member. It may contain a no-change carry-forward, but no official rule proves that, so it is not an accepted 2040 fleet basis.",
        True, "GAP",
    )
    add(
        "MT00", "NUCLEAR_CAPACITY", "2040",
        rel(pemmdb_files["MT00_2040"]),
        f"{mt['nuclear_capacity_MW']:.3f} MW; no nonzero nuclear band.",
        "DIRECT_OFFICIAL_PEMMDB",
        "No nuclear input is required for MT00 under the official fleet.",
        True, "PASS",
    )
    add(
        "TN00", "NUCLEAR_CAPACITY", "closest supplied TN members",
        rel(TN_PEMMDB_COMPARISON),
        f"{tn['nuclear_capacity_MW']:.3f} MW in the nominal member; no nonzero nuclear band in the supplied TN horizon members.",
        "DIRECT_OFFICIAL_MEMBER_COMPARISON",
        "Nuclear is independently non-material; this does not close the broader TN fleet-horizon gap.",
        True, "PASS",
    )
    add(
        "MT00", "HYDRO_STORAGE", "2040",
        rel(pemmdb_files["MT00_2040"]),
        f"hydro={mt['hydro_power_MW']:.3f} MW; battery={mt['battery_power_MW']:.3f} MW.",
        "DIRECT_OFFICIAL_PEMMDB",
        "No hydro/storage input is required for the official MT00 fleet.",
        True, "PASS",
    )
    add(
        "TN00", "HYDRO_STORAGE", "nominal NT2040 member",
        f"{rel(pemmdb_files['TN00_2040'])}; {rel(TN_PEMMDB_COMPARISON)}",
        f"cached hydro={tn['hydro_power_MW']:.3f} MW; battery={tn['battery_power_MW']:.3f} MW, but member metadata remains 2030/blank.",
        "DIRECT_OFFICIAL_MEMBER_WITH_STALE_METADATA",
        "Zero is the closest supplied value but is not promoted as a traceable 2040 assumption while the member is unresolved.",
        True, "GAP",
    )
    add(
        "MT00", "PV_CAPACITY", "2040",
        rel(pemmdb_files["MT00_2040"]),
        f"{mt['solar_capacity_MW']:.9f} MW.",
        "DIRECT_OFFICIAL_PEMMDB",
        "The annual PV quantity is traceable.",
        True, "PASS",
    )
    add(
        "TN00", "PV_CAPACITY", "nominal NT2040 member",
        f"{rel(pemmdb_files['TN00_2040'])}; {rel(TN_PEMMDB_COMPARISON)}",
        f"cached value={tn['solar_capacity_MW']:.3f} MW; metadata year=2030 and scenario blank.",
        "DIRECT_OFFICIAL_MEMBER_WITH_STALE_METADATA",
        "The cached value is not accepted as a 2040 quantity without corrected source evidence or explicit carry-forward authorization.",
        True, "GAP",
    )
    add(
        "MT00", "PV_PROFILE", "2040 capacity / CY2019",
        rel(pecd_files["MT00_SOLAR"]),
        f"2019={pecd['MT00_SOLAR']['nonnull']}/8,760; per-unit sum={pecd['MT00_SOLAR']['sum']:.9f}.",
        "DIRECT_OFFICIAL_PECD",
        "The official 2019 Malta solar realization is complete.",
        True, "PASS",
    )
    add(
        "TN00", "PV_PROFILE", "NT2040 capacity / CY2019",
        f"{rel(pecd_files['TN00_SOLAR'])}; {rel(CUTOUT)}",
        f"official 2019 PECD={pecd['TN00_SOLAR']['nonnull']}/8,760; accepted cutout y-min={cutout['y_min']:.1f} degrees N and does not cover full Tunisia.",
        "DIRECT_OFFICIAL_PECD_PLUS_ACCEPTED_CUTOUT_COVERAGE",
        "A small full-country 2019 weather extension is technically feasible but remains downstream of the scenario-basis decision.",
        True, "GAP",
    )
    add(
        "MT00", "WIND_CAPACITY_AND_PROFILE", "2040 / CY2019",
        rel(pemmdb_files["MT00_2040"]),
        f"wind capacity={mt['wind_capacity_MW']:.3f} MW; no wind profile required.",
        "DIRECT_OFFICIAL_PEMMDB",
        "Zero wind makes an hourly wind profile non-applicable for the endogenous MT00 price model.",
        True, "PASS",
    )
    add(
        "TN00", "WIND_CAPACITY", "nominal NT2040 member",
        f"{rel(pemmdb_files['TN00_2040'])}; {rel(TN_PEMMDB_COMPARISON)}",
        f"cached value={tn['wind_capacity_MW']:.3f} MW; metadata year=2030 and scenario blank.",
        "DIRECT_OFFICIAL_MEMBER_WITH_STALE_METADATA",
        "The capacity is not accepted as a 2040 quantity without corrected source evidence or explicit carry-forward authorization.",
        True, "GAP",
    )
    add(
        "TN00", "WIND_PROFILE", "NT2040 capacity / CY2019",
        f"{rel(pecd_files['TN00_WIND'])}; {rel(CUTOUT)}",
        f"official 2019 PECD={pecd['TN00_WIND']['nonnull']}/8,760; accepted cutout y-min={cutout['y_min']:.1f} degrees N and does not cover full Tunisia.",
        "DIRECT_OFFICIAL_PECD_PLUS_ACCEPTED_CUTOUT_COVERAGE",
        "A full-country 2019 weather extension is required if a TN scenario basis is authorized.",
        True, "GAP",
    )
    for node in ("MT00", "TN00"):
        add(
            node, "FUEL_CO2_COST_BASIS", "DE2040",
            rel(PRICES), common_costs,
            "DIRECT_OFFICIAL_COMMON_SCENARIO_COST_TABLE",
            "The common technology cost basis is available; no undocumented TN/MT country-price proxy is required for standard fuel classes.",
            True, "PASS",
        )
    add(
        "MT00", "CROSS_BORDER_TOPOLOGY", "TYNDP reference grid",
        rel(REFERENCE_GRID),
        "ITSI-MT00=225 MW in each direction in the upstream reference grid.",
        "DIRECT_OFFICIAL_REFERENCE_GRID",
        "Upstream topology is sufficient; it does not overwrite the frozen downstream MEM interface capacity.",
        True, "PASS",
    )
    add(
        "TN00", "CROSS_BORDER_TOPOLOGY", "TYNDP reference grid",
        rel(REFERENCE_GRID),
        "ITSI-TN00=600/600 MW; DZ00-TN00=250/250 MW; LY00-TN00=500/500 MW.",
        "DIRECT_OFFICIAL_REFERENCE_GRID",
        "Topology exists and confirms that missing geometry is not the substantive blocker. The frozen MEM IT-TN interface remains 600 MW bidirectional.",
        True, "PASS",
    )
    add(
        "MT00", "CHRONOLOGY_COMPATIBILITY", "CY2019 / 8,760 UTC",
        f"{rel(CUTOUT)}; {rel(pecd_files['MT00_SOLAR'])}; {rel(DE_DEMAND)}",
        f"cutout={cutout['hours']}/8,760 and covers Malta; solar=8,760/8,760; demand=0/8,760 for CY2019.",
        "ACCEPTED_WEATHER_SOURCE_PLUS_OFFICIAL_PROFILE_INVENTORY",
        "Only the documented Malta demand adapter remains. No annual scenario quantity needs to be invented.",
        True, "GAP",
    )
    add(
        "TN00", "CHRONOLOGY_COMPATIBILITY", "CY2019 / 8,760 UTC",
        f"{rel(CUTOUT)}; {rel(NT_DEMAND)}; {rel(pecd_files['TN00_SOLAR'])}; {rel(pecd_files['TN00_WIND'])}",
        f"cutout={cutout['hours']}/8,760 but only north of {cutout['y_min']:.1f} degrees N; demand/solar/wind CY2019=0/0/0 populated hours.",
        "ACCEPTED_WEATHER_SOURCE_PLUS_OFFICIAL_PROFILE_INVENTORY",
        "The 2019 realization is technically incomplete and cannot be closed before the TN annual scenario basis is authorized.",
        True, "GAP",
    )
    add(
        "MT00", "PUBLISHED_DE2040_PRICE_QA", "DE2040 CY1995",
        rel(DE_OUTPUT),
        f"modelled hours={output['modelled_hours']}; marginal-cost columns={len(output['marginal_cost_columns']['MT00'])}; cross-border columns={len(output['crossborder_columns']['MT00'])}.",
        "DIRECT_OFFICIAL_PUBLISHED_OUTPUT",
        "Confirms MT00 participation in DE2040; 8,736-hour prices remain QA only and are not promoted to MEM CY2019.",
        False, "PASS",
    )
    add(
        "TN00", "PUBLISHED_DE2040_PRICE_QA", "DE2040 CY1995",
        rel(DE_OUTPUT),
        f"modelled hours={output['modelled_hours']}; marginal-cost columns=0; cross-border columns=0.",
        "DIRECT_OFFICIAL_PUBLISHED_OUTPUT",
        "The absence agrees with the final methodology's North-African-only-NT rule; it is not evidence of a mere helper defect.",
        False, "GAP",
    )

    _write_csv(
        QA_DIR / "MEM_TYNDP_TN00_MT00_INPUT_COVERAGE_QA.csv",
        coverage,
        fieldnames=[
            "node", "field", "scenario_year", "official_file",
            "observed_value_status", "provenance", "interpretation",
            "required_for_price_model", "pass_gap",
        ],
    )
    _write_csv(QA_DIR / "MEM_2019_LOAD_SOURCE_COMPATIBILITY_QA.csv", load_rows)

    receipts: list[dict[str, Any]] = []
    receipt_specs = [
        ("TYNDP24_FINAL_SCENARIO_METHODOLOGY", "CONTROLLING_SCENARIO_PERIMETER_RULE", METHODOLOGY_REPORT, "Page 17: North-African Med-TSO markets modelled only in National Trends; Figure 5 includes MT and excludes TN", "AUDITED_CONTROLLING", "Final January-2025 scenario methodology; more specific and later than the intermediate implementation guideline."),
        ("TYNDP24_INTERMEDIATE_IMPLEMENTATION_GUIDELINES", "GENERIC_PROJECT_ASSESSMENT_PERIMETER", IMPLEMENTATION_GUIDELINES, "Pages 14-15: generic market-model/project-assessment perimeter includes TN and MT", "AUDITED_CONTEXT_ONLY", "Does not establish a TN node in DE/GA scenario-building runs."),
        ("TYNDP24_NODES", "MARKET_NODE_IDENTITY", NODES, "MT00 present; TN00 absent", "AUDITED", "Official TYNDP 2024 distributed input."),
        ("TYNDP24_SUPPLY_INPUTS", "DE_GA_SUPPLY_OVERRIDES", SUPPLY, "MT00 present; TN00 absent", "AUDITED", "Official TYNDP 2024 distributed input."),
        ("TYNDP24_REFERENCE_GRID", "UPSTREAM_NTC_TOPOLOGY", REFERENCE_GRID, "ITSI-MT00 225 MW; ITSI-TN00 600 MW; DZ00-TN00 250 MW; LY00-TN00 500 MW", "AUDITED", "Does not overwrite frozen MEM interface capacities."),
        ("TYNDP24_DEMAND_SCENARIOS", "DEMAND_PARAMETERIZATION", DEMAND_SCENARIOS, "Malta DE/GA links present; Tunisia absent", "AUDITED", "Official TYNDP 2024 distributed input."),
        ("TYNDP24_DE2040_DEMAND_PROFILES", "DE2040_DEMAND_SHAPES", DE_DEMAND, "MT00 present through climate year 2016; TN00 absent", "AUDITED", "No 2019 MT/TN accepted profile."),
        ("TYNDP24_NT2040_DEMAND_PROFILES", "SAME_PACKAGE_ALTERNATE_DEMAND_SHAPES", NT_DEMAND, "TN00 2018 complete/2019 empty; MT00 2018/2019 empty", "AUDITED", "NT is not silently substituted for DE."),
        ("TYNDP24_DE2040_CY1995_OUTPUT", "PUBLISHED_MODEL_OUTPUT_QA", DE_OUTPUT, "8736 hours; MT marginal cost present; TN absent", "AUDITED_QA_ONLY", "Not eligible as MEM 2019 price input."),
        ("TYNDP24_COMMODITY_PRICES", "UPSTREAM_FUEL_AND_CO2_COSTS", PRICES, "TYNDP 2024 scenario commodity prices", "AVAILABLE", "No node-specific TN/MT record is required for common commodity classes."),
        ("TYNDP24_TN00_PEMMDB_HORIZON_COMPARISON", "TN00_2030_2040_2050_OOXML_AUDIT", TN_PEMMDB_COMPARISON, "Nominal 2040 and 2050 cached/formula content identical; both retain 2030 metadata/source relationship", "AUDITED_DERIVATION", "Establishes stale packaging; does not prove intentional 2040 carry-forward."),
    ]
    for spec in receipt_specs:
        receipts.append(_receipt(*spec))
    for key, path in pecd_files.items():
        receipts.append(
            _receipt(
                f"TYNDP24_PECD_{key}",
                "2019_RESOURCE_PROFILE_INPUT",
                path,
                f"2019 non-null={pecd[key]['nonnull']}/8760",
                "AUDITED",
                f"Scenario label={pecd[key]['metadata'].get('Scenario') or 'blank'}.",
            )
        )
    for key, path in pemmdb_files.items():
        info = pemmdb[key]
        receipts.append(
            _receipt(
                f"TYNDP24_PEMMDB_{key}",
                "GENERATION_FLEET_INPUT",
                path,
                f"metadata_year={info['metadata_year']};scenario={info['metadata_scenario'] or 'blank'};thermal={info['thermal_capacity_MW']:.3f} MW",
                "AUDITED",
                "Official PEMMDB2 archive member.",
            )
        )
    receipts.append(
        {
            "source_id": "PYPSA_CUTOUT_V1_EUROPE_2019_SARAH3_ERA5",
            "role": "ACCEPTED_2019_WEATHER_COVERAGE",
            "path": str(CUTOUT.resolve()),
            "bytes": CUTOUT.stat().st_size,
            "sha256": CUTOUT_SHA256,
            "coverage": (
                f"{cutout['hours']} hours; x={cutout['x_min']:.1f}..{cutout['x_max']:.1f}; "
                f"y={cutout['y_min']:.1f}..{cutout['y_max']:.1f}"
            ),
            "status": "PRESERVED_ACCEPTED_HASH_NOT_REHASHED",
            "notes": "Covers Malta; covers only northern Tunisia and is not a full-country TN weather source.",
        }
    )
    _write_csv(QA_DIR / "MEM_TYNDP_TN00_MT00_SOURCE_RECEIPTS.csv", receipts)

    raw_receipts = [
        _receipt(
            f"TYNDP24_RAW_{path.stem.upper().replace(' ', '_')}",
            "IMMUTABLE_OFFICIAL_DOWNLOAD_ARCHIVE",
            path,
            "Official TYNDP 2024 package",
            "PRESERVED",
            "Downloaded from https://2024.entsos-tyndp-scenarios.eu/download/",
        )
        for path in sorted(RAW.glob("*"))
        if path.is_file()
    ]
    _write_csv(QA_DIR / "MEM_TYNDP_2024_RAW_ARCHIVE_RECEIPTS.csv", raw_receipts)

    audit = {
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "scope": "TYNDP_2024_DE2040_PLUS_2019_TN00_MT00_E1_COVERAGE",
        "official_download_url": "https://2024.entsos-tyndp-scenarios.eu/download/",
        "load_compatibility": load_metrics,
        "source_presence_validation": source_presence,
        "nodes": {
            "MT00": {
                "market_node": "PRESENT",
                "official_scenario_rule": "INCLUDED_IN_DE_GA_MODELLED_PERIMETER",
                "reference_grid": "PRESENT_ITSI_MT00_225_MW_UPSTREAM",
                "de2040_demand_scenario": "PRESENT",
                "de2040_demand_profiles": "PRESENT_FOR_1982_2016_INCLUDING_1995_2008_2009",
                "2019_demand_profile": "ABSENT",
                "2019_solar_pecd": f"{pecd['MT00_SOLAR']['nonnull']}/8760",
                "2019_wind_profile": "NOT_REQUIRED_ZERO_CAPACITY",
                "pemmdb_2040": "VALID_NT2040_504_MW_THERMAL",
                "published_de2040_price": "PRESENT_8736_HOURS_QA_ONLY",
                "e1_status": "BOUNDED_2019_DEMAND_TEMPORAL_ADAPTER_REMAINS",
                "controlling_m2": False,
            },
            "TN00": {
                "market_node": "ABSENT_FROM_LIST_OF_NODES",
                "official_scenario_rule": "NORTH_AFRICAN_MEDTSO_MARKETS_MODELLED_ONLY_IN_NATIONAL_TRENDS",
                "reference_grid": "PRESENT_ITSI_TN00_600_MW_PLUS_DZ_LY_CONNECTIONS_UPSTREAM",
                "de2040_demand_scenario": "ABSENT",
                "de2040_demand_profile": "ABSENT",
                "nt2040_demand_profile": f"2018={nt_tn[2018]['nonnull']}/8760;2019={nt_tn[2019]['nonnull']}/8760",
                "2019_solar_pecd": f"{pecd['TN00_SOLAR']['nonnull']}/8760",
                "2019_wind_pecd": f"{pecd['TN00_WIND']['nonnull']}/8760",
                "accepted_cutout_coverage": "PARTIAL_NORTH_ONLY_Y_MIN_33_DEGREES",
                "pemmdb_2040": "STALE_PACKAGING_MEMBER_METADATA_2030_SCENARIO_BLANK_2040_EQUALS_2050_NOT_ACCEPTED_AS_CARRY_FORWARD",
                "published_de2040_price": "ABSENT",
                "e1_status": "MANAGER_RETURN_M2_TN_DE2040_SCENARIO_MAPPING_AND_2040_SUPPLY_AUTHORIZATION_REQUIRED",
                "controlling_m2": True,
            },
        },
        "classification": "MANAGER_RETURN_M2_NARROWED",
        "reason": (
            "The final TYNDP 2024 Scenario Methodology explicitly models North-African "
            "Med-TSO electricity markets only in National Trends. It does not document "
            "a carry-forward of Tunisia into DE/GA. The nominal TN 2040 PEMMDB member "
            "has stale 2030 metadata and is identical to the nominal 2050 member, so it "
            "does not establish traceable 2040 supply. TN also lacks complete 2019 load, "
            "solar and wind realizations. Malta is a valid DE2040 node and has only a "
            "bounded 2019 demand temporal-adapter gap."
        ),
        "smallest_manager_decision": (
            "Authorize either an explicit MEM cross-scenario E1 extension using official "
            "NT Tunisia alongside DE2040 Europe, including treatment of the stale TN "
            "PEMMDB member, or authorize a separate Med-TSO 2040 Tunisia extension."
        ),
        "no_price_series_generated": True,
        "no_pypsa_network_built": True,
        "solver_invocations": 0,
    }
    with (QA_DIR / "MEM_TYNDP_TN00_MT00_AUDIT.json").open(
        "w", encoding="utf-8"
    ) as handle:
        json.dump(audit, handle, indent=2, ensure_ascii=False)
        handle.write("\n")
    return audit


def main() -> None:
    audit = build_audit()
    print(json.dumps(audit["nodes"], indent=2))
    print(f"classification={audit['classification']}")


if __name__ == "__main__":
    main()

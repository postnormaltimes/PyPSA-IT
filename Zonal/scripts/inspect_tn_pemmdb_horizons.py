"""Compare official TN00 PEMMDB workbooks at 2030, 2040, and 2050.

This is a source audit. It does not transform or promote any PEMMDB value into
an MEM runtime input.
"""

from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
EXTRACTED = ROOT / "runtime_sources" / "tyndp_2024" / "extracted"
FILES = {
    2030: EXTRACTED
    / "PEMMDB2_compare/PEMMDB2/2030/PEMMDB_TN00_NationalTrends_2030.xlsx",
    2040: EXTRACTED
    / "PEMMDB2/PEMMDB2/2040/PEMMDB_TN00_NationalTrends_2040.xlsx",
    2050: EXTRACTED
    / "PEMMDB2_compare/PEMMDB2/2050/PEMMDB_TN00_NationalTrends_2050.xlsx",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def snapshot(path: Path, *, data_only: bool) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(
        path,
        read_only=False,
        data_only=data_only,
        keep_links=True,
    )
    try:
        return {
            worksheet.title: {
                cell.coordinate: cell.value
                for row in worksheet.iter_rows()
                for cell in row
                if cell.value is not None
            }
            for worksheet in workbook.worksheets
        }
    finally:
        workbook.close()


def package_links(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with zipfile.ZipFile(path) as package:
        for name in package.namelist():
            if "externalLink" not in name and not name.endswith(
                ("workbook.xml", "workbook.xml.rels")
            ):
                continue
            text = package.read(name).decode("utf-8", "ignore")
            rows.append(
                {
                    "member": name,
                    "targets": re.findall(r'Target="([^"]+)"', text),
                    "values": re.findall(r'(?:val|sheetName)="([^"]+)"', text),
                }
            )
    return rows


def cell_differences(
    left: dict[str, dict[str, Any]], right: dict[str, dict[str, Any]]
) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for sheet in sorted(set(left) | set(right)):
        left_cells = left.get(sheet, {})
        right_cells = right.get(sheet, {})
        for coordinate in sorted(set(left_cells) | set(right_cells)):
            left_value = left_cells.get(coordinate)
            right_value = right_cells.get(coordinate)
            if left_value != right_value:
                output.append(
                    {
                        "sheet": sheet,
                        "coordinate": coordinate,
                        "left": left_value,
                        "right": right_value,
                    }
                )
    return output


def main() -> None:
    missing = [str(path) for path in FILES.values() if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing PEMMDB members: " + "; ".join(missing))

    cached = {year: snapshot(path, data_only=True) for year, path in FILES.items()}
    formulas = {year: snapshot(path, data_only=False) for year, path in FILES.items()}
    report: dict[str, Any] = {"files": {}, "comparisons": {}}

    for year, path in FILES.items():
        formula_cells = [
            {"sheet": sheet, "coordinate": coordinate, "formula": value}
            for sheet, cells in formulas[year].items()
            for coordinate, value in cells.items()
            if isinstance(value, str) and value.startswith("=")
        ]
        thermal = cached[year].get("Thermal", {})
        report["files"][str(year)] = {
            "path": str(path.resolve()),
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
            "sheet_names": list(cached[year]),
            "thermal_metadata": {
                coordinate: thermal.get(coordinate)
                for coordinate in ("B3", "B4", "B5", "B6", "B7")
            },
            "formula_cell_count": len(formula_cells),
            "formula_cells": formula_cells,
            "package_links": package_links(path),
        }

    for left, right in ((2030, 2040), (2040, 2050), (2030, 2050)):
        cached_diff = cell_differences(cached[left], cached[right])
        formula_diff = cell_differences(formulas[left], formulas[right])
        report["comparisons"][f"{left}_vs_{right}"] = {
            "cached_visible_cell_difference_count": len(cached_diff),
            "cached_visible_cell_differences": cached_diff,
            "formula_cell_difference_count": len(formula_diff),
            "formula_cell_differences": formula_diff,
        }

    output = (
        ROOT
        / "qa"
        / "runtime_closure"
        / "MEM_TYNDP_TN00_PEMMDB_2030_2040_2050_COMPARISON.json"
    )
    output.write_text(
        json.dumps(report, indent=2, ensure_ascii=False, default=str) + "\n",
        encoding="utf-8",
    )
    print(output)
    for pair, comparison in report["comparisons"].items():
        print(
            pair,
            "cached_differences=",
            comparison["cached_visible_cell_difference_count"],
            "formula_differences=",
            comparison["formula_cell_difference_count"],
        )


if __name__ == "__main__":
    main()

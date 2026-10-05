"""Deterministic ETX-7B5 fixed Stage-A topology contract builder.

The module emits machine-readable node, interface, directional-capacity,
authority, provenance, carry-forward, and exclusion contracts.  It never
imports PyPSA, constructs a production network, configures a solver, or runs
optimization.
"""

from __future__ import annotations

import argparse
import hashlib
import itertools
import json
import os
from pathlib import Path
from typing import Any

import pandas as pd
import yaml
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
CONFIG_PATH = ROOT / "config" / "stage_a_topology.yaml"
DEFAULT_OUTPUT_DIR = ROOT / "stage_a_inputs" / "topology_v1_0"
SCHEMA_VERSION = "ETX7B5_TOPOLOGY_V1_0"
MARKETS = ("AT", "CH", "FR", "GR", "HR", "IT", "ME", "MT", "SI", "TN")
HORIZONS = (2040, 2050)

OUTPUT_FILES = {
    "nodes": "MEM_ETX7B5_Stage_A_Node_Registry_v1.0.csv",
    "interfaces": "MEM_ETX7B5_Physical_Interface_Registry_v1.0.csv",
    "links": "MEM_ETX7B5_Directional_Link_Registry_v1.0.csv",
    "capacities_2040": "MEM_ETX7B5_Fixed_Directional_Capacities_2040_v1.0.csv",
    "capacities_2050": "MEM_ETX7B5_Fixed_Directional_Capacities_2050_v1.0.csv",
    "authority_matrix": "MEM_ETX7B5_Topology_Authority_Matrix_v1.0.csv",
    "provenance": "MEM_ETX7B5_Topology_Provenance_v1.0.csv",
    "carry_forward": "MEM_ETX7B5_2040_to_2050_Carry_Forward_Dispositions_v1.0.csv",
    "no_links": "MEM_ETX7B5_Explicit_No_Link_Controls_v1.0.csv",
    "excluded_markets": "MEM_ETX7B5_Excluded_Market_Controls_v1.0.csv",
    "manifest": "MEM_ETX7B5_Topology_Output_Manifest_v1.0.csv",
}

MARKET_NAMES = {
    "AT": "Austria",
    "CH": "Switzerland",
    "FR": "France",
    "GR": "Greece",
    "HR": "Croatia",
    "IT": "Italy",
    "ME": "Montenegro",
    "MT": "Malta",
    "SI": "Slovenia",
    "TN": "Tunisia",
}


def _io_name(path: Path) -> str:
    """Return a Windows long-path-safe name without changing canonical paths."""
    resolved = str(path.resolve())
    if os.name == "nt" and not resolved.startswith("\\\\?\\") and len(resolved) >= 240:
        return "\\\\?\\" + resolved
    return resolved


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with open(_io_name(path), "rb") as handle:
        while chunk := handle.read(chunk_size):
            digest.update(chunk)
    return digest.hexdigest().upper()


def _stat(path: Path) -> os.stat_result:
    return os.stat(_io_name(path))


def _exists(path: Path) -> bool:
    return os.path.exists(_io_name(path))


def _resolve(value: str | Path) -> Path:
    candidate = Path(value)
    return candidate if candidate.is_absolute() else (ROOT / candidate).resolve()


def _relative(path: Path) -> str:
    resolved = path.resolve()
    return resolved.relative_to(ROOT).as_posix() if resolved.is_relative_to(ROOT) else resolved.as_posix()


def _write_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    frame.to_csv(path, index=False, encoding="utf-8", lineterminator="\n", na_rep="")


def _read_csv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, dtype=str, keep_default_na=False)


def _row_count(path: Path) -> int:
    return len(_read_csv(path))


def load_config(path: Path = CONFIG_PATH) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        return yaml.safe_load(handle)


def _authority_path(authority: dict[str, Any], key: str = "path") -> Path:
    return _resolve(authority[key])


def _validate_authority_hashes(config: dict[str, Any]) -> None:
    for authority in config["authorities"].values():
        path = _authority_path(authority)
        if not _exists(path) or sha256_file(path) != authority["sha256"]:
            raise ValueError(f"Topology authority hash mismatch: {_relative(path)}")
        if authority.get("raw_archive_path"):
            raw = _resolve(authority["raw_archive_path"])
            if not _exists(raw) or sha256_file(raw) != authority["raw_archive_sha256"]:
                raise ValueError(f"Topology raw-archive hash mismatch: {_relative(raw)}")


def _validate_mem_interfaces(config: dict[str, Any]) -> None:
    authority = config["authorities"]["mem_italy_interface"]
    source = _read_csv(_authority_path(authority))
    source.columns = [column.strip().lower() for column in source.columns]
    expected_markets = {record["external_market"] for record in config["italy_interfaces"]}
    if expected_markets != {"FR", "CH", "AT", "SI", "ME", "GR", "MT", "TN"}:
        raise ValueError("Italy-facing interface scope is not the accepted eight-market set")
    for record in config["italy_interfaces"]:
        market = record["external_market"]
        for direction, key in (("IMPORT", "external_to_it_MW"), ("EXPORT", "it_to_external_MW")):
            selected = source.loc[
                source["external_market"].eq(market) & source["direction"].eq(direction)
            ]
            if len(selected) != 1:
                raise ValueError(f"Frozen MEM interface row not unique: {market} {direction}")
            observed = float(selected.iloc[0]["capacity_mw"])
            if observed != float(record[key]):
                raise ValueError(f"Frozen MEM interface capacity mismatch: {market} {direction}")
            if selected.iloc[0]["italian_zone"] != record["stage_b_receiving_zone"]:
                raise ValueError(f"Stage-B receiving-zone mismatch: {market}")
            if selected.iloc[0]["applicable_years"] != "2040|2050":
                raise ValueError(f"Frozen MEM interface horizon mismatch: {market}")


def _validate_tyndp_connections(config: dict[str, Any]) -> None:
    authority = config["authorities"]["tyndp_horizon_construction"]
    workbook = load_workbook(
        _io_name(_authority_path(authority)), read_only=True, data_only=True
    )
    try:
        reference = workbook[authority["reference_sheet"]]
        investments = workbook[authority["investment_sheet"]]
        for record in config["external_connections"]:
            ref = tuple(
                reference.cell(row=int(record["reference_row"]), column=column).value
                for column in range(1, 4)
            )
            expected_ref = (
                record["source_border"],
                record["reference_a_to_b_MW"],
                record["reference_b_to_a_MW"],
            )
            if ref != expected_ref:
                raise ValueError(f"TYNDP reference-grid row mismatch: {record['physical_id']}")

            accepted: list[tuple[Any, ...]] = []
            for project in record["accepted_real_increments_through_2040"]:
                row = int(project["investment_row"])
                values = tuple(investments.cell(row=row, column=column).value for column in range(1, 8))
                expected = (
                    project["project"],
                    record["source_border"].split("-")[0],
                    record["source_border"].split("-")[1],
                    project["commissioning_year"],
                    "All",
                    project["a_to_b_increment_MW"],
                    project["b_to_a_increment_MW"],
                )
                if values != expected or " Real " not in f" {values[0]} ":
                    raise ValueError(f"TYNDP accepted real-project row mismatch: {project['project']}")
                accepted.append((values[0], row, values[3], values[5], values[6]))

            source_a, source_b = record["source_border"].split("-")
            observed_real: list[tuple[Any, ...]] = []
            for row_number, values in enumerate(investments.iter_rows(min_row=2, values_only=True), start=2):
                name, from_node, to_node, year, _, direct, indirect = values[:7]
                if (
                    from_node == source_a
                    and to_node == source_b
                    and isinstance(name, str)
                    and " Real " in f" {name} "
                    and int(year) <= 2040
                ):
                    observed_real.append((name, row_number, year, direct, indirect))
            if sorted(accepted) != sorted(observed_real):
                raise ValueError(f"TYNDP real-project coverage mismatch: {record['physical_id']}")

            expected_a_to_b = float(record["reference_a_to_b_MW"]) + sum(
                float(item["a_to_b_increment_MW"])
                for item in record["accepted_real_increments_through_2040"]
            )
            expected_b_to_a = float(record["reference_b_to_a_MW"]) + sum(
                float(item["b_to_a_increment_MW"])
                for item in record["accepted_real_increments_through_2040"]
            )
            if expected_a_to_b != float(record["capacity_2040_a_to_b_MW"]):
                raise ValueError(f"2040 a-to-b construction mismatch: {record['physical_id']}")
            if expected_b_to_a != float(record["capacity_2040_b_to_a_MW"]):
                raise ValueError(f"2040 b-to-a construction mismatch: {record['physical_id']}")
    finally:
        workbook.close()


def validate_config(config: dict[str, Any]) -> None:
    if tuple(config["scope"]["markets"]) != MARKETS:
        raise ValueError("Stage-A topology must contain the exact canonical ten-market order")
    if tuple(int(year) for year in config["scope"]["horizons"]) != HORIZONS:
        raise ValueError("Stage-A topology must contain exactly 2040 and 2050")
    if bool(config["scope"]["p_nom_extendable"]):
        raise ValueError("Stage-A link expansion is forbidden")
    if config["scope"]["corsica"] != "EXCLUDED_STAGE_A_STAGE_B_REGULATED_ZERO_INJECTION_ONLY":
        raise ValueError("CORS exclusion is not locked")
    if config["scope"]["krsko_physical_host"] != "SI":
        raise ValueError("Krsko physical host must remain SI")

    physical_ids = [record["physical_id"] for record in config["italy_interfaces"]]
    physical_ids += [record["physical_id"] for record in config["external_connections"]]
    if len(physical_ids) != len(set(physical_ids)):
        raise ValueError("Duplicate physical interface identity in topology config")
    pairs = [{"IT", record["external_market"]} for record in config["italy_interfaces"]]
    pairs += [
        {record["endpoint_a"], record["endpoint_b"]}
        for record in config["external_connections"]
    ]
    pair_keys = [tuple(sorted(pair)) for pair in pairs]
    if len(pair_keys) != 12 or len(set(pair_keys)) != 12:
        raise ValueError("Topology config must contain exactly 12 unique physical connections")
    if any(not pair <= set(MARKETS) for pair in pairs):
        raise ValueError("Topology config contains a market outside the frozen perimeter")

    _validate_authority_hashes(config)
    _validate_mem_interfaces(config)
    _validate_tyndp_connections(config)


def build_node_registry(config: dict[str, Any]) -> pd.DataFrame:
    stage_b_zone = {
        record["external_market"]: record["stage_b_receiving_zone"]
        for record in config["italy_interfaces"]
    }
    rows = []
    for market in MARKETS:
        rows.append(
            {
                "node_id": f"STAGE_A_{market}",
                "market": market,
                "market_name": MARKET_NAMES[market],
                "node_granularity": "ONE_NODE_PER_MARKET",
                "scenario": "Base",
                "horizons": "2040|2050",
                "capacity_expansion_allowed": False,
                "stage_b_price_receiving_zone": stage_b_zone.get(market, ""),
                "stage_b_boundary_role": (
                    "DIRECT_EXTERNAL_PRICE_BOUNDARY"
                    if market in stage_b_zone
                    else "ITALY_AGGREGATE"
                    if market == "IT"
                    else "STAGE_A_ONLY_NO_DIRECT_STAGE_B_PRICE_BOUNDARY"
                ),
                "krsko_physical_nuclear_host": market == "SI",
                "croatia_krsko_entitlement_role": (
                    "QA_ACCOUNTING_ONLY_NO_PHYSICAL_ASSET" if market == "HR" else "NOT_APPLICABLE"
                ),
                "status": "ETX7B5_FIXED_STAGE_A_NODE",
            }
        )
    return pd.DataFrame.from_records(rows)


def _link_row(
    *,
    physical_id: str,
    family: str,
    from_market: str,
    to_market: str,
    year: int,
    capacity: float,
    stage_b_zone: str,
    authority_id: str,
    authority_classification: str,
    authority_horizon: str,
    implementation_rule: str,
    direct_vs_carry_forward: str,
    evidence_class: str,
    provenance_path: str,
    provenance_location: str,
    status: str,
) -> dict[str, Any]:
    return {
        "physical_link_id": physical_id,
        "directional_link_id": f"ETX7B5_{year}_{from_market}_TO_{to_market}",
        "link_family": family,
        "from_market": from_market,
        "to_market": to_market,
        "year": year,
        "scenario": "Base",
        "capacity_MW": capacity,
        "unit": "MW",
        "direction_convention": "POSITIVE_FROM_MARKET_TO_TO_MARKET",
        "stage_b_receiving_zone": stage_b_zone,
        "authority_id": authority_id,
        "authority_classification": authority_classification,
        "authority_horizon": authority_horizon,
        "implementation_rule": implementation_rule,
        "direct_vs_carry_forward": direct_vs_carry_forward,
        "accepted_evidence_class": evidence_class,
        "p_nom_extendable": False,
        "fixed_capacity": True,
        "interconnector_asset_type": "NETWORK_LINK_NOT_GENERATOR",
        "provenance_path": provenance_path,
        "provenance_location": provenance_location,
        "status": status,
    }


def build_directional_links(config: dict[str, Any]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    mem_authority = config["authorities"]["mem_italy_interface"]
    for interface in config["italy_interfaces"]:
        for year in HORIZONS:
            rows.append(
                _link_row(
                    physical_id=interface["physical_id"],
                    family="ITALY_FACING",
                    from_market=interface["external_market"],
                    to_market="IT",
                    year=year,
                    capacity=interface["external_to_it_MW"],
                    stage_b_zone=interface["stage_b_receiving_zone"],
                    authority_id=mem_authority["source_id"],
                    authority_classification="FROZEN_MEM_INTERFACE_CONTRACT",
                    authority_horizon=config["horizon_rules"]["italy_interface_authority_horizon"],
                    implementation_rule="FROZEN_MEM_IMPORT_EXTERNAL_TO_IT",
                    direct_vs_carry_forward=config["horizon_rules"]["italy_interface_direct_vs_carry_forward"],
                    evidence_class="ACCEPTED_FROZEN_CONTRACT_A",
                    provenance_path=mem_authority["path"],
                    provenance_location=f"CSV rows {interface['source_rows']}",
                    status=config["horizon_rules"]["italy_interface_status"],
                )
            )
            rows.append(
                _link_row(
                    physical_id=interface["physical_id"],
                    family="ITALY_FACING",
                    from_market="IT",
                    to_market=interface["external_market"],
                    year=year,
                    capacity=interface["it_to_external_MW"],
                    stage_b_zone=interface["stage_b_receiving_zone"],
                    authority_id=mem_authority["source_id"],
                    authority_classification="FROZEN_MEM_INTERFACE_CONTRACT",
                    authority_horizon=config["horizon_rules"]["italy_interface_authority_horizon"],
                    implementation_rule="FROZEN_MEM_EXPORT_IT_TO_EXTERNAL",
                    direct_vs_carry_forward=config["horizon_rules"]["italy_interface_direct_vs_carry_forward"],
                    evidence_class="ACCEPTED_FROZEN_CONTRACT_A",
                    provenance_path=mem_authority["path"],
                    provenance_location=f"CSV rows {interface['source_rows']}",
                    status=config["horizon_rules"]["italy_interface_status"],
                )
            )

    tyndp = config["authorities"]["tyndp_horizon_construction"]
    for connection in config["external_connections"]:
        for year in HORIZONS:
            rule = config["horizon_rules"][str(year)]
            if year == 2040:
                implementation = (
                    "REFERENCE_GRID_BASELINE_PLUS_ACCEPTED_REAL_PROJECT_INCREMENTS_"
                    "COMMISSIONED_BY_2040;CONCEPT_CANDIDATES_EXCLUDED"
                )
            else:
                implementation = (
                    "EXPLICIT_2040_NTC_CARRY_FORWARD_TO_2050;"
                    "NO_INTERPOLATION_OR_EXTRAPOLATION"
                )
            source_rows = [f"reference row {connection['reference_row']}"]
            source_rows += [
                f"investment row {item['investment_row']}"
                for item in connection["accepted_real_increments_through_2040"]
            ]
            for from_market, to_market, capacity_key in (
                (connection["endpoint_a"], connection["endpoint_b"], "capacity_2040_a_to_b_MW"),
                (connection["endpoint_b"], connection["endpoint_a"], "capacity_2040_b_to_a_MW"),
            ):
                rows.append(
                    _link_row(
                        physical_id=connection["physical_id"],
                        family="EXTERNAL_EXTERNAL",
                        from_market=from_market,
                        to_market=to_market,
                        year=year,
                        capacity=connection[capacity_key],
                        stage_b_zone="",
                        authority_id=tyndp["source_id"],
                        authority_classification="OFFICIAL_TYNDP24_ACCEPTED_COMPOSITE",
                        authority_horizon=str(rule["external_external_authority_horizon"]),
                        implementation_rule=implementation,
                        direct_vs_carry_forward=rule["external_external_direct_vs_carry_forward"],
                        evidence_class="ACCEPTED_OFFICIAL_COMPOSITE_A",
                        provenance_path=tyndp["path"],
                        provenance_location=(
                            f"{tyndp['reference_sheet']} and {tyndp['investment_sheet']}; "
                            + ", ".join(source_rows)
                        ),
                        status=rule["external_external_status"],
                    )
                )

    frame = pd.DataFrame.from_records(rows)
    asymmetric = frame.groupby(["physical_link_id", "year"])["capacity_MW"].transform("nunique").gt(1)
    frame["asymmetry_treatment"] = asymmetric.map(
        {
            True: "PRESERVE_EXPLICIT_DIRECTIONAL_ASYMMETRY",
            False: "EXPLICIT_SYMMETRIC_CAPACITY_TWO_DIRECTED_ROWS",
        }
    )
    frame["reverse_directional_link_id"] = frame.apply(
        lambda row: f"ETX7B5_{row['year']}_{row['to_market']}_TO_{row['from_market']}", axis=1
    )
    return frame.sort_values(
        ["year", "link_family", "physical_link_id", "from_market", "to_market"],
        kind="stable",
    ).reset_index(drop=True)


def build_interface_registry(links: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for physical_id, group in links.groupby("physical_link_id", sort=True):
        endpoints = sorted(set(group["from_market"]) | set(group["to_market"]))
        row: dict[str, Any] = {
            "physical_link_id": physical_id,
            "endpoint_a": endpoints[0],
            "endpoint_b": endpoints[1],
            "link_family": group["link_family"].iloc[0],
            "stage_b_receiving_zone": "|".join(sorted(set(group["stage_b_receiving_zone"]) - {""})),
            "directional_rows_2040": int(group["year"].eq(2040).sum()),
            "directional_rows_2050": int(group["year"].eq(2050).sum()),
            "asymmetric_2040": group.loc[group["year"].eq(2040), "capacity_MW"].nunique() > 1,
            "asymmetric_2050": group.loc[group["year"].eq(2050), "capacity_MW"].nunique() > 1,
            "capacity_expansion_allowed": False,
            "asset_representation": "TWO_EXPLICIT_DIRECTED_FIXED_LINK_ROWS_PER_HORIZON",
            "status": "PHYSICAL_INTERFACE_IDENTITY_COMPLETE",
        }
        rows.append(row)
    return pd.DataFrame.from_records(rows).sort_values("physical_link_id", kind="stable").reset_index(drop=True)


def build_capacity_table(links: pd.DataFrame, year: int) -> pd.DataFrame:
    columns = [
        "directional_link_id",
        "physical_link_id",
        "link_family",
        "from_market",
        "to_market",
        "year",
        "capacity_MW",
        "unit",
        "direction_convention",
        "p_nom_extendable",
        "fixed_capacity",
        "asymmetry_treatment",
        "status",
    ]
    return links.loc[links["year"].eq(year), columns].reset_index(drop=True)


def build_authority_matrix(links: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "physical_link_id",
        "directional_link_id",
        "link_family",
        "from_market",
        "to_market",
        "year",
        "capacity_MW",
        "unit",
        "authority_id",
        "authority_classification",
        "authority_horizon",
        "implementation_rule",
        "direct_vs_carry_forward",
        "asymmetry_treatment",
        "accepted_evidence_class",
        "provenance_path",
        "provenance_location",
        "status",
    ]
    frame = links[columns].copy()
    return frame.rename(columns={"authority_id": "authority"})


def build_provenance(config: dict[str, Any], config_path: Path) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []

    def add(source_id: str, path: Path, role: str, horizon: str, transformation: str) -> None:
        rows.append(
            {
                "source_id": source_id,
                "path": _relative(path),
                "bytes": _stat(path).st_size,
                "sha256": sha256_file(path),
                "accepted_role": role,
                "authority_horizon": horizon,
                "transformation_boundary": transformation,
                "status": "ACCEPTED_AUTHORITY_LOCKED",
            }
        )

    plan = config["controlling_plan"]
    add("MEM_REMAINING_EXECUTION_PLAN_V2_2", _resolve(plan["path"]), "CONTROLLING_B5_PLAN", "2040|2050", "NO_SCOPE_EXPANSION")
    add("ETX7B5_TOPOLOGY_CONFIG", config_path.resolve(), "B5_IMPLEMENTATION_AUTHORITY_LOCK", "2040|2050", "CONFIG_TO_DETERMINISTIC_CONTRACTS")
    for authority in config["authorities"].values():
        add(authority["source_id"], _authority_path(authority), authority["role"], "2040|2050", "SOURCE_ROWS_TO_EXPLICIT_DIRECTED_CAPACITY")
        if authority.get("raw_archive_path"):
            add("TYNDP24_RAW_GRID_INVESTMENT_ARCHIVE", _resolve(authority["raw_archive_path"]), "IMMUTABLE_OFFICIAL_DOWNLOAD_ARCHIVE", "SOURCE_PACKAGE", "HASH_RECEIPT_ONLY")
    for name, lock in config["predecessor_locks"].items():
        add(f"PREDECESSOR_LOCK_{name.upper()}", _resolve(lock["path"]), "IMMUTABLE_PREDECESSOR_RECEIPT", "PRE_B5", "HASH_REGRESSION_ONLY")
    return pd.DataFrame.from_records(rows).drop_duplicates("source_id").sort_values("source_id", kind="stable").reset_index(drop=True)


def build_carry_forward(links: pd.DataFrame) -> pd.DataFrame:
    selected = links.loc[
        links["link_family"].eq("EXTERNAL_EXTERNAL") & links["year"].eq(2050)
    ].copy()
    return pd.DataFrame(
        {
            "physical_link_id": selected["physical_link_id"],
            "directional_link_id_2050": selected["directional_link_id"],
            "from_market": selected["from_market"],
            "to_market": selected["to_market"],
            "source_year": 2040,
            "target_year": 2050,
            "source_capacity_MW": selected["capacity_MW"],
            "target_capacity_MW": selected["capacity_MW"],
            "unit": "MW",
            "rule": "2040_NTC_CARRY_FORWARD_TO_2050",
            "interpolation_used": False,
            "extrapolation_used": False,
            "status": "ACCEPTED_2040_CARRY_FORWARD_TO_2050",
        }
    ).reset_index(drop=True)


def build_no_link_controls(config: dict[str, Any]) -> pd.DataFrame:
    physical_pairs = {
        tuple(sorted(("IT", record["external_market"])))
        for record in config["italy_interfaces"]
    }
    physical_pairs |= {
        tuple(sorted((record["endpoint_a"], record["endpoint_b"])))
        for record in config["external_connections"]
    }
    rows = []
    for endpoint_a, endpoint_b in itertools.combinations(MARKETS, 2):
        if (endpoint_a, endpoint_b) in physical_pairs:
            continue
        rows.append(
            {
                "control_id": f"NO_LINK_{endpoint_a}_{endpoint_b}",
                "endpoint_a": endpoint_a,
                "endpoint_b": endpoint_b,
                "year_scope": "2040|2050",
                "capacity_MW": 0,
                "unit": "MW",
                "disposition": "NO_LINK_WITHIN_STAGE_A_PERIMETER",
                "hidden_transit_node_allowed": False,
                "reason": "NOT_AN_ACCEPTED_RETAINED_MARKET_PHYSICAL_CONNECTION",
                "status": "EXPLICIT_NO_LINK_CONTROL",
            }
        )
    return pd.DataFrame.from_records(rows)


def build_excluded_market_controls(config: dict[str, Any]) -> pd.DataFrame:
    return pd.DataFrame.from_records(
        [
            {
                "market": market,
                "stage_a_bus_allowed": False,
                "stage_a_link_endpoint_allowed": False,
                "disposition": "EXCLUDED_FROM_FIXED_STAGE_A_PERIMETER",
                "status": "EXPLICIT_EXCLUSION_CONTROL",
            }
            for market in config["excluded_markets"]
        ]
    )


def build_topology_frames(config: dict[str, Any], config_path: Path = CONFIG_PATH) -> dict[str, pd.DataFrame]:
    nodes = build_node_registry(config)
    links = build_directional_links(config)
    return {
        "nodes": nodes,
        "interfaces": build_interface_registry(links),
        "links": links,
        "capacities_2040": build_capacity_table(links, 2040),
        "capacities_2050": build_capacity_table(links, 2050),
        "authority_matrix": build_authority_matrix(links),
        "provenance": build_provenance(config, config_path),
        "carry_forward": build_carry_forward(links),
        "no_links": build_no_link_controls(config),
        "excluded_markets": build_excluded_market_controls(config),
    }


def _manifest(output_dir: Path) -> pd.DataFrame:
    rows = []
    for key, filename in OUTPUT_FILES.items():
        if key == "manifest":
            continue
        path = output_dir / filename
        rows.append(
            {
                "artifact": key,
                "relative_path": f"stage_a_inputs/topology_v1_0/{filename}",
                "bytes": path.stat().st_size,
                "rows": _row_count(path),
                "sha256": sha256_file(path),
                "schema_version": SCHEMA_VERSION,
                "status": "DETERMINISTIC_B5_OUTPUT",
            }
        )
    return pd.DataFrame.from_records(rows).sort_values("artifact", kind="stable").reset_index(drop=True)


def build_topology_package(
    output_dir: Path = DEFAULT_OUTPUT_DIR,
    *,
    config_path: Path = CONFIG_PATH,
) -> dict[str, Any]:
    config_path = config_path.resolve()
    config = load_config(config_path)
    validate_config(config)
    output_dir.mkdir(parents=True, exist_ok=True)
    frames = build_topology_frames(config, config_path)
    for key, frame in frames.items():
        _write_csv(frame, output_dir / OUTPUT_FILES[key])
    manifest = _manifest(output_dir)
    _write_csv(manifest, output_dir / OUTPUT_FILES["manifest"])
    frames["manifest"] = manifest
    return {"frames": frames, "output_dir": output_dir}


def main() -> None:
    parser = argparse.ArgumentParser(description="Build deterministic ETX-7B5 Stage-A topology contracts")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--config", type=Path, default=CONFIG_PATH)
    args = parser.parse_args()
    result = build_topology_package(args.output_dir, config_path=args.config)
    print(
        json.dumps(
            {
                "status": "BUILT_NOT_VALIDATED",
                "output_dir": str(result["output_dir"]),
                "counts": {key: len(value) for key, value in result["frames"].items()},
                "production_network": "NOT_CONSTRUCTED",
                "optimization": "NOT_RUN",
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
FROZEN_SOURCE_DIR = ROOT / "runtime_sources" / "etx7a" / "frozen"
WORKBOOK_NAME = "MEM_ETX7A_External_Country_Harmonised_Freeze_v1.0.xlsx"
STATE_OF_RECORD_NAME = (
    "MEM_ETX7A_External_Country_Harmonised_Freeze_State_of_Record_v1.0.md"
)
SOURCE_MANIFEST_NAME = "MEM_ETX7A_FROZEN_SOURCE_MANIFEST.csv"
ETX7A_VERSION = "1.0"

CANONICAL_SHEET = "02_CANONICAL_STACK"
RUNTIME_SHEET = "03_RUNTIME_INPUTS"
HEADER_ROW = 4

RUNTIME_ROLES = (
    "IMPLEMENT",
    "IMPLEMENT_TOTAL",
    "IMPLEMENT_STORAGE_PAIR",
    "IMPLEMENT_CARRIER_SPLIT",
)
NON_RUNTIME_ROLES = ("DERIVATION_DETAIL_ONLY", "QA_CONTROL_ONLY")
EXPECTED_ROLE_COUNTS = {
    "IMPLEMENT": 98,
    "IMPLEMENT_TOTAL": 3,
    "IMPLEMENT_STORAGE_PAIR": 48,
    "IMPLEMENT_CARRIER_SPLIT": 12,
}
COUNTRIES = ("FR", "CH", "AT", "SI", "HR", "ME", "GR", "MT", "TN")
HORIZONS = (2040, 2050)


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(chunk_size):
            digest.update(chunk)
    return digest.hexdigest()


def _normalise_header(value: object) -> str:
    text = str(value).strip().lower()
    text = text.replace("2040 ", "value_2040_", 1) if text.startswith("2040 ") else text
    text = text.replace("2050 ", "value_2050_", 1) if text.startswith("2050 ") else text
    text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")
    aliases = {
        "value_2040_value": "value_2040",
        "value_2050_value": "value_2050",
        "value_2040_evidence": "evidence_2040",
        "value_2050_evidence": "evidence_2050",
        "value_2040_quality": "quality_2040",
        "value_2050_quality": "quality_2050",
        "value_2040_source_id": "source_id_2040",
        "value_2050_source_id": "source_id_2050",
        "value_2040_closure": "closure_2040",
        "value_2050_closure": "closure_2050",
    }
    return aliases.get(text, text)


def _read_table(sheet_name: str, header_row: int = HEADER_ROW) -> pd.DataFrame:
    workbook_path = FROZEN_SOURCE_DIR / WORKBOOK_NAME
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        headers = [_normalise_header(cell.value) for cell in worksheet[header_row]]
        records: list[dict[str, object]] = []
        for workbook_row, values in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            if not any(value is not None for value in values):
                continue
            record = dict(zip(headers, values, strict=False))
            record["etx7a_sheet"] = sheet_name
            record["etx7a_workbook_row"] = workbook_row
            records.append(record)
    finally:
        workbook.close()
    return pd.DataFrame.from_records(records)


def read_canonical_evidence() -> pd.DataFrame:
    """Read all 207 ETX-7A canonical evidence rows at original source grain."""

    return _read_table(CANONICAL_SHEET)


def read_runtime_inputs() -> pd.DataFrame:
    """Read and enrich the 161 runtime-authoritative ETX-7A rows."""

    runtime = _read_table(RUNTIME_SHEET)
    unexpected = sorted(set(runtime["runtime_role"].dropna()) - set(RUNTIME_ROLES))
    if unexpected:
        raise ValueError(f"Unexpected ETX-7A runtime roles: {unexpected}")

    canonical = read_canonical_evidence()
    keys = [
        "country_code",
        "category",
        "source_group",
        "harmonised_group",
        "unit",
        "runtime_role",
        "source_workbook",
        "source_row",
    ]
    lookup = canonical.loc[
        canonical["runtime_role"].isin(RUNTIME_ROLES),
        keys + ["source_sheet", "etx7a_workbook_row"],
    ].rename(columns={"etx7a_workbook_row": "canonical_etx7a_workbook_row"})
    if lookup.duplicated(keys).any():
        duplicates = lookup.loc[lookup.duplicated(keys, keep=False), keys]
        raise ValueError(
            "ETX-7A canonical-to-runtime provenance key is not unique: "
            f"{duplicates.to_dict(orient='records')}"
        )
    runtime = runtime.merge(lookup, on=keys, how="left", validate="one_to_one")
    if runtime["canonical_etx7a_workbook_row"].isna().any():
        missing = runtime.loc[
            runtime["canonical_etx7a_workbook_row"].isna(), keys
        ].to_dict(orient="records")
        raise ValueError(f"Runtime rows without canonical provenance: {missing}")
    runtime["canonical_etx7a_workbook_row"] = runtime[
        "canonical_etx7a_workbook_row"
    ].astype(int)
    return runtime


def read_frozen_source_manifest() -> pd.DataFrame:
    return pd.read_csv(FROZEN_SOURCE_DIR / SOURCE_MANIFEST_NAME, dtype=str)


def verify_frozen_sources() -> pd.DataFrame:
    """Return one deterministic receipt row for each frozen ETX-7A source."""

    manifest = read_frozen_source_manifest()
    receipts: list[dict[str, object]] = []
    for row in manifest.to_dict(orient="records"):
        path = ROOT / Path(str(row["relative_path"]))
        observed_bytes = path.stat().st_size if path.exists() else None
        observed_hash = sha256_file(path).upper() if path.exists() else None
        expected_bytes = int(str(row["bytes"]))
        expected_hash = str(row["sha256"]).upper()
        status = (
            "PASS"
            if path.exists()
            and observed_bytes == expected_bytes
            and observed_hash == expected_hash
            else "FAIL"
        )
        receipts.append(
            {
                "relative_path": row["relative_path"],
                "expected_bytes": expected_bytes,
                "observed_bytes": observed_bytes,
                "expected_sha256": expected_hash,
                "observed_sha256": observed_hash,
                "role": row["role"],
                "status": status,
            }
        )
    return pd.DataFrame.from_records(receipts)


def assert_columns(frame: pd.DataFrame, required: Iterable[str], label: str) -> None:
    missing = sorted(set(required) - set(frame.columns))
    if missing:
        raise ValueError(f"{label} is missing required columns: {missing}")
